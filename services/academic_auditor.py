from __future__ import annotations

import re
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn


CHAPTER_PATTERN = re.compile(
    r"^\s*BAB\s+([IVXLCDM]+|\d+)\b(?:\s*[:.\-]?\s*(.*))?$",
    re.IGNORECASE,
)

SUBHEADING_PATTERN = re.compile(
    r"^\s*(\d+(?:\.\d+){1,3})\s+(.+)$"
)

TABLE_CAPTION_PATTERN = re.compile(
    r"^\s*Tabel\s+(\d+)(?:\.(\d+))?\s+(.+)$",
    re.IGNORECASE,
)

FIGURE_CAPTION_PATTERN = re.compile(
    r"^\s*(?:Gambar|Figure)\s+(\d+)(?:\.(\d+))?\s+(.+)$",
    re.IGNORECASE,
)

APPENDIX_PATTERN = re.compile(
    r"^\s*Lampiran(?:\s+(\d+|[A-Z]))?\b",
    re.IGNORECASE,
)

YEAR_PATTERN = re.compile(
    r"\b(19|20)\d{2}[a-z]?\b",
    re.IGNORECASE,
)

PARENTHETICAL_CITATION_PATTERN = re.compile(
    r"\(([^()]{1,100}?),\s*((?:19|20)\d{2}[a-z]?)\)",
    re.IGNORECASE,
)

NARRATIVE_CITATION_PATTERN = re.compile(
    r"\b([A-ZÀ-ÖØ-Ý][A-Za-zÀ-ÖØ-öø-ÿ'’\-]{2,})"
    r"(?:\s+et\s+al\.)?\s*\("
    r"((?:19|20)\d{2}[a-z]?)\)",
)

NUMERIC_CITATION_PATTERN = re.compile(
    r"\[(\d+(?:\s*[-,]\s*\d+)*)\]"
)

FRONT_MATTER_PATTERNS = {
    "halaman_judul": re.compile(
        r"^\s*(HALAMAN\s+JUDUL|JUDUL)\s*$",
        re.IGNORECASE,
    ),
    "persetujuan": re.compile(
        r"^\s*(LEMBAR|HALAMAN)\s+PERSETUJUAN\s*$",
        re.IGNORECASE,
    ),
    "pengesahan": re.compile(
        r"^\s*(LEMBAR|HALAMAN)\s+PENGESAHAN\s*$",
        re.IGNORECASE,
    ),
    "keaslian": re.compile(
        r"^\s*(SURAT\s+)?PERNYATAAN\s+KEASLIAN.*$",
        re.IGNORECASE,
    ),
    "abstrak": re.compile(
        r"^\s*ABSTRAK\s*$",
        re.IGNORECASE,
    ),
    "abstract": re.compile(
        r"^\s*ABSTRACT\s*$",
        re.IGNORECASE,
    ),
    "kata_pengantar": re.compile(
        r"^\s*KATA\s+PENGANTAR\s*$",
        re.IGNORECASE,
    ),
    "daftar_isi": re.compile(
        r"^\s*DAFTAR\s+ISI\s*$",
        re.IGNORECASE,
    ),
    "daftar_tabel": re.compile(
        r"^\s*DAFTAR\s+TABEL\s*$",
        re.IGNORECASE,
    ),
    "daftar_gambar": re.compile(
        r"^\s*DAFTAR\s+GAMBAR\s*$",
        re.IGNORECASE,
    ),
    "daftar_lampiran": re.compile(
        r"^\s*DAFTAR\s+LAMPIRAN\s*$",
        re.IGNORECASE,
    ),
    "daftar_pustaka": re.compile(
        r"^\s*(DAFTAR\s+PUSTAKA|REFERENCES?)\s*$",
        re.IGNORECASE,
    ),
}


DEFAULT_POLICIES: dict[str, dict[str, Any]] = {
    "skripsi": {
        "font": "Times New Roman",
        "font_size": 12.0,
        "line_spacing": 2.0,
        "margin_top_cm": 4.0,
        "margin_bottom_cm": 3.0,
        "margin_left_cm": 4.0,
        "margin_right_cm": 3.0,
        "required_front_matter": [
            "abstrak",
            "kata_pengantar",
            "daftar_isi",
        ],
        "required_chapters": [1, 2, 3, 4, 5],
        "require_references": True,
    },
    "laporan": {
        "font": "Arial",
        "font_size": 11.0,
        "line_spacing": 1.5,
        "margin_top_cm": 3.0,
        "margin_bottom_cm": 3.0,
        "margin_left_cm": 3.0,
        "margin_right_cm": 3.0,
        "required_front_matter": [
            "daftar_isi",
        ],
        "required_chapters": [1, 2, 3, 4, 5],
        "require_references": True,
    },
    "jurnal": {
        "font": "Times New Roman",
        "font_size": 11.0,
        "line_spacing": 1.15,
        "margin_top_cm": 2.5,
        "margin_bottom_cm": 2.5,
        "margin_left_cm": 2.5,
        "margin_right_cm": 2.5,
        "required_front_matter": [
            "abstrak",
        ],
        "required_chapters": [],
        "require_references": True,
    },
}


ROMAN_VALUES = {
    "I": 1,
    "V": 5,
    "X": 10,
    "L": 50,
    "C": 100,
    "D": 500,
    "M": 1000,
}


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def clean_text(value: str) -> str:
    return re.sub(
        r"\s+",
        " ",
        value or "",
    ).strip()


def roman_to_int(value: str) -> int | None:
    normalized = value.upper().strip()

    if normalized.isdigit():
        return int(normalized)

    total = 0
    previous = 0

    for character in reversed(normalized):
        number = ROMAN_VALUES.get(character)

        if number is None:
            return None

        if number < previous:
            total -= number
        else:
            total += number
            previous = number

    return total or None


def safe_cm(value: Any) -> float | None:
    try:
        return round(value.cm, 2)
    except (
        AttributeError,
        TypeError,
    ):
        return None


def safe_pt(value: Any) -> float | None:
    try:
        return round(value.pt, 2)
    except (
        AttributeError,
        TypeError,
    ):
        return None


def get_run_font_name(run: Any) -> str | None:
    if run.font.name:
        return run.font.name

    try:
        properties = run._element.rPr

        if properties is None:
            return None

        fonts = properties.rFonts

        if fonts is None:
            return None

        return (
            fonts.get(qn("w:ascii"))
            or fonts.get(qn("w:hAnsi"))
            or fonts.get(qn("w:eastAsia"))
        )
    except (
        AttributeError,
        TypeError,
    ):
        return None


def contains_field(
    document: Document,
    field_name: str,
) -> bool:
    instructions = document._element.xpath(
        ".//w:instrText"
    )

    return any(
        field_name.upper()
        in clean_text(node.text or "").upper()
        for node in instructions
    )


def add_finding(
    findings: list[dict[str, Any]],
    severity: str,
    code: str,
    title: str,
    message: str,
    location: str = "",
    recommendation: str = "",
) -> None:
    findings.append(
        {
            "severity": severity,
            "code": code,
            "title": title,
            "message": message,
            "location": location,
            "recommendation": recommendation,
        }
    )


def detect_front_matter(
    paragraphs: list[str],
) -> dict[str, bool]:
    result = {
        key: False
        for key in FRONT_MATTER_PATTERNS
    }

    for text in paragraphs:
        for key, pattern in FRONT_MATTER_PATTERNS.items():
            if pattern.match(text):
                result[key] = True

    return result


def extract_chapters(
    paragraph_records: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    chapters: list[dict[str, Any]] = []

    for record in paragraph_records:
        match = CHAPTER_PATTERN.match(
            record["text"]
        )

        if not match:
            continue

        raw_number = match.group(1)
        number = roman_to_int(raw_number)

        chapters.append(
            {
                "paragraph_index": record["index"],
                "raw_number": raw_number,
                "number": number,
                "title": clean_text(
                    match.group(2) or ""
                ),
                "text": record["text"],
                "style": record["style"],
            }
        )

    return chapters


def extract_subheadings(
    paragraph_records: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    subheadings: list[dict[str, Any]] = []

    for record in paragraph_records:
        match = SUBHEADING_PATTERN.match(
            record["text"]
        )

        if not match:
            continue

        numbering = match.group(1)
        segments = [
            int(segment)
            for segment in numbering.split(".")
        ]

        subheadings.append(
            {
                "paragraph_index": record["index"],
                "numbering": numbering,
                "segments": segments,
                "title": match.group(2),
                "text": record["text"],
                "style": record["style"],
                "level": len(segments),
            }
        )

    return subheadings


def extract_captions(
    paragraph_records: list[dict[str, Any]],
    pattern: re.Pattern[str],
) -> list[dict[str, Any]]:
    captions: list[dict[str, Any]] = []

    for record in paragraph_records:
        match = pattern.match(
            record["text"]
        )

        if not match:
            continue

        chapter_number = int(match.group(1))
        sequence_value = match.group(2)

        sequence_number = (
            int(sequence_value)
            if sequence_value
            else chapter_number
        )

        if sequence_value:
            full_number = (
                f"{chapter_number}.{sequence_number}"
            )
        else:
            full_number = str(chapter_number)

        title_group = (
            match.group(3)
            if match.lastindex
            and match.lastindex >= 3
            else ""
        )

        captions.append(
            {
                "paragraph_index": record["index"],
                "chapter": chapter_number,
                "sequence": sequence_number,
                "numbering": full_number,
                "title": clean_text(title_group),
                "text": record["text"],
                "style": record["style"],
            }
        )

    return captions


def audit_chapter_sequence(
    chapters: list[dict[str, Any]],
    findings: list[dict[str, Any]],
    required_chapters: list[int],
) -> list[dict[str, Any]]:
    checklist: list[dict[str, Any]] = []

    observed = [
        chapter["number"]
        for chapter in chapters
        if chapter["number"] is not None
    ]

    duplicates = [
        number
        for number, count in Counter(
            observed
        ).items()
        if count > 1
    ]

    if duplicates:
        add_finding(
            findings,
            "error",
            "duplicate_chapter_number",
            "Nomor BAB ganda",
            (
                "Nomor BAB berikut ditemukan lebih dari "
                f"satu kali: {duplicates}."
            ),
            recommendation=(
                "Periksa penomoran judul BAB dan hapus "
                "duplikasi."
            ),
        )

    if observed:
        expected = list(
            range(
                min(observed),
                max(observed) + 1,
            )
        )

        if observed != expected:
            add_finding(
                findings,
                "error",
                "chapter_sequence_inconsistent",
                "Urutan BAB tidak konsisten",
                (
                    f"Urutan terdeteksi {observed}, "
                    f"sedangkan urutan kontinu {expected}."
                ),
                recommendation=(
                    "Susun kembali BAB secara berurutan."
                ),
            )

    for required_number in required_chapters:
        exists = required_number in observed

        checklist.append(
            {
                "item": f"BAB {required_number}",
                "status": (
                    "passed"
                    if exists
                    else "failed"
                ),
                "message": (
                    "Terdeteksi"
                    if exists
                    else "Belum terdeteksi"
                ),
            }
        )

        if not exists:
            add_finding(
                findings,
                "warning",
                "required_chapter_missing",
                f"BAB {required_number} belum terdeteksi",
                (
                    f"Dokumen belum memiliki judul "
                    f"BAB {required_number} yang dapat dikenali."
                ),
                recommendation=(
                    "Pastikan judul BAB ditulis dengan format "
                    "'BAB I', 'BAB II', dan seterusnya."
                ),
            )

    for chapter in chapters:
        if not chapter["style"].lower().startswith(
            "heading 1"
        ):
            add_finding(
                findings,
                "warning",
                "chapter_heading_style",
                "Style judul BAB belum konsisten",
                (
                    f"'{chapter['text']}' menggunakan style "
                    f"'{chapter['style']}', bukan Heading 1."
                ),
                location=(
                    f"Paragraf {chapter['paragraph_index'] + 1}"
                ),
                recommendation=(
                    "Gunakan style Heading 1 pada setiap "
                    "judul BAB."
                ),
            )

    return checklist


def audit_subheading_sequence(
    subheadings: list[dict[str, Any]],
    findings: list[dict[str, Any]],
) -> None:
    seen_numbers: set[str] = set()
    grouped: dict[
        tuple[int, ...],
        list[dict[str, Any]],
    ] = defaultdict(list)

    for item in subheadings:
        numbering = item["numbering"]

        if numbering in seen_numbers:
            add_finding(
                findings,
                "error",
                "duplicate_subheading_number",
                "Nomor subbab ganda",
                (
                    f"Nomor subbab {numbering} ditemukan "
                    "lebih dari satu kali."
                ),
                location=(
                    f"Paragraf {item['paragraph_index'] + 1}"
                ),
                recommendation=(
                    "Perbaiki nomor subbab yang berulang."
                ),
            )

        seen_numbers.add(numbering)

        segments = item["segments"]

        if len(segments) >= 2:
            parent = tuple(
                segments[:-1]
            )
            grouped[parent].append(item)

        expected_style = (
            "heading 2"
            if item["level"] == 2
            else "heading 3"
        )

        if not item["style"].lower().startswith(
            expected_style
        ):
            add_finding(
                findings,
                "warning",
                "subheading_style",
                "Style subbab belum sesuai",
                (
                    f"'{item['text']}' menggunakan "
                    f"style '{item['style']}'."
                ),
                location=(
                    f"Paragraf {item['paragraph_index'] + 1}"
                ),
                recommendation=(
                    "Gunakan Heading 2 untuk subbab dan "
                    "Heading 3 untuk sub-subbab."
                ),
            )

    for parent, items in grouped.items():
        observed = [
            item["segments"][-1]
            for item in items
        ]

        expected = list(
            range(
                1,
                len(observed) + 1,
            )
        )

        if observed != expected:
            parent_text = ".".join(
                str(value)
                for value in parent
            )

            add_finding(
                findings,
                "warning",
                "subheading_sequence_gap",
                "Urutan subbab tidak kontinu",
                (
                    f"Pada kelompok {parent_text}, "
                    f"urutan terdeteksi {observed}; "
                    f"seharusnya {expected}."
                ),
                recommendation=(
                    "Periksa nomor subbab yang hilang, "
                    "tertukar, atau berulang."
                ),
            )


def audit_caption_sequence(
    captions: list[dict[str, Any]],
    caption_type: str,
    findings: list[dict[str, Any]],
) -> None:
    numbers = [
        caption["numbering"]
        for caption in captions
    ]

    duplicates = [
        value
        for value, count in Counter(
            numbers
        ).items()
        if count > 1
    ]

    if duplicates:
        add_finding(
            findings,
            "error",
            f"duplicate_{caption_type}_caption",
            f"Nomor caption {caption_type} ganda",
            (
                "Nomor caption berikut ditemukan lebih "
                f"dari satu kali: {duplicates}."
            ),
            recommendation=(
                f"Perbaiki urutan nomor caption "
                f"{caption_type}."
            ),
        )

    grouped: dict[int, list[int]] = defaultdict(list)

    for caption in captions:
        grouped[
            caption["chapter"]
        ].append(
            caption["sequence"]
        )

        if caption["style"].lower() != "caption":
            add_finding(
                findings,
                "info",
                f"{caption_type}_caption_style",
                f"Style caption {caption_type}",
                (
                    f"'{caption['text']}' belum menggunakan "
                    "style Caption Microsoft Word."
                ),
                location=(
                    f"Paragraf {caption['paragraph_index'] + 1}"
                ),
                recommendation=(
                    "Gunakan style Caption agar daftar otomatis "
                    "dapat dibuat lebih stabil."
                ),
            )

    for chapter, sequences in grouped.items():
        expected = list(
            range(
                1,
                len(sequences) + 1,
            )
        )

        if sequences != expected:
            add_finding(
                findings,
                "warning",
                f"{caption_type}_caption_sequence",
                f"Urutan caption {caption_type} tidak kontinu",
                (
                    f"Pada BAB {chapter}, urutan terdeteksi "
                    f"{sequences}; seharusnya {expected}."
                ),
                recommendation=(
                    "Periksa caption yang hilang atau "
                    "penomorannya tidak tepat."
                ),
            )


def extract_reference_entries(
    paragraph_records: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    references: list[dict[str, Any]] = []
    inside_references = False

    for record in paragraph_records:
        text = record["text"]

        if FRONT_MATTER_PATTERNS[
            "daftar_pustaka"
        ].match(text):
            inside_references = True
            continue

        if inside_references and CHAPTER_PATTERN.match(text):
            break

        if inside_references and APPENDIX_PATTERN.match(text):
            break

        if not inside_references:
            continue

        if len(text) < 5:
            continue

        year_match = YEAR_PATTERN.search(text)
        first_token_match = re.match(
            r"^\s*([A-Za-zÀ-ÖØ-öø-ÿ'’\-]+)",
            text,
        )

        references.append(
            {
                "paragraph_index": record["index"],
                "text": text,
                "year": (
                    year_match.group(0).lower()
                    if year_match
                    else None
                ),
                "author_key": (
                    first_token_match.group(1).lower()
                    if first_token_match
                    else None
                ),
            }
        )

    return references


def extract_citations(
    paragraph_records: list[dict[str, Any]],
    reference_start_index: int | None,
) -> list[dict[str, Any]]:
    citations: list[dict[str, Any]] = []

    for record in paragraph_records:
        if (
            reference_start_index is not None
            and record["index"] >= reference_start_index
        ):
            continue

        text = record["text"]

        for match in PARENTHETICAL_CITATION_PATTERN.finditer(
            text
        ):
            author_text = clean_text(
                match.group(1)
            )

            author_key = re.split(
                r"\s+|&|dan|;",
                author_text,
                maxsplit=1,
                flags=re.IGNORECASE,
            )[0].strip(
                " ,.;:"
            ).lower()

            citations.append(
                {
                    "type": "author_year",
                    "author_key": author_key,
                    "year": match.group(2).lower(),
                    "text": match.group(0),
                    "paragraph_index": record["index"],
                }
            )

        for match in NARRATIVE_CITATION_PATTERN.finditer(
            text
        ):
            citations.append(
                {
                    "type": "author_year",
                    "author_key": match.group(1).lower(),
                    "year": match.group(2).lower(),
                    "text": match.group(0),
                    "paragraph_index": record["index"],
                }
            )

        for match in NUMERIC_CITATION_PATTERN.finditer(
            text
        ):
            citations.append(
                {
                    "type": "numeric",
                    "author_key": None,
                    "year": None,
                    "text": match.group(0),
                    "paragraph_index": record["index"],
                }
            )

    return citations


def audit_citations_and_references(
    citations: list[dict[str, Any]],
    references: list[dict[str, Any]],
    findings: list[dict[str, Any]],
    require_references: bool,
) -> dict[str, Any]:
    reference_keys = {
        (
            reference["author_key"],
            reference["year"],
        )
        for reference in references
        if reference["author_key"]
        and reference["year"]
    }

    citation_keys = {
        (
            citation["author_key"],
            citation["year"],
        )
        for citation in citations
        if citation["type"] == "author_year"
        and citation["author_key"]
        and citation["year"]
    }

    missing_reference_keys = sorted(
        citation_keys - reference_keys
    )

    uncited_reference_keys = sorted(
        reference_keys - citation_keys
    )

    if require_references and not references:
        add_finding(
            findings,
            "error",
            "references_missing",
            "Daftar pustaka belum terdeteksi",
            (
                "Bagian daftar pustaka atau entri referensi "
                "belum dapat dikenali."
            ),
            recommendation=(
                "Tambahkan bagian Daftar Pustaka dan pastikan "
                "setiap sumber ditulis pada paragraf terpisah."
            ),
        )

    for author_key, year in missing_reference_keys[:30]:
        add_finding(
            findings,
            "warning",
            "citation_without_reference",
            "Sitasi belum ditemukan dalam daftar pustaka",
            (
                f"Sitasi indikatif '{author_key.title()}, "
                f"{year}' tidak memiliki pasangan yang "
                "terdeteksi dalam daftar pustaka."
            ),
            recommendation=(
                "Periksa nama penulis dan tahun pada sitasi "
                "serta daftar pustaka."
            ),
        )

    for author_key, year in uncited_reference_keys[:30]:
        add_finding(
            findings,
            "info",
            "reference_not_cited",
            "Referensi belum ditemukan dalam teks",
            (
                f"Referensi indikatif '{author_key.title()}, "
                f"{year}' belum memiliki pasangan sitasi "
                "yang terdeteksi."
            ),
            recommendation=(
                "Pastikan referensi memang digunakan atau "
                "hapus sumber yang tidak dirujuk."
            ),
        )

    missing_year_references = [
        reference
        for reference in references
        if not reference["year"]
    ]

    for reference in missing_year_references[:20]:
        add_finding(
            findings,
            "warning",
            "reference_year_missing",
            "Tahun referensi belum terdeteksi",
            (
                f"Entri berikut tidak memiliki tahun yang "
                f"terdeteksi: {reference['text'][:120]}"
            ),
            location=(
                f"Paragraf {reference['paragraph_index'] + 1}"
            ),
            recommendation=(
                "Periksa kembali format tahun pada entri "
                "daftar pustaka."
            ),
        )

    return {
        "citations_total": len(citations),
        "references_total": len(references),
        "citation_keys_total": len(citation_keys),
        "reference_keys_total": len(reference_keys),
        "citations_without_reference": len(
            missing_reference_keys
        ),
        "references_not_cited": len(
            uncited_reference_keys
        ),
        "numeric_citations_total": sum(
            1
            for citation in citations
            if citation["type"] == "numeric"
        ),
        "note": (
            "Pencocokan sitasi bersifat indikatif karena "
            "variasi gaya penulisan nama penulis."
        ),
    }


def audit_formatting(
    document: Document,
    paragraph_records: list[dict[str, Any]],
    policy: dict[str, Any],
    findings: list[dict[str, Any]],
) -> dict[str, Any]:
    font_names: Counter[str] = Counter()
    font_sizes: Counter[str] = Counter()
    line_spacings: Counter[str] = Counter()
    alignments: Counter[str] = Counter()

    total_runs = 0
    total_body_paragraphs = 0
    expected_font = str(
        policy.get("font") or ""
    ).lower()

    expected_size = float(
        policy.get("font_size", 12)
    )

    expected_spacing = float(
        policy.get("line_spacing", 2)
    )

    wrong_font_runs = 0
    wrong_size_runs = 0

    for record in paragraph_records:
        paragraph = record["paragraph"]

        if (
            CHAPTER_PATTERN.match(record["text"])
            or SUBHEADING_PATTERN.match(record["text"])
            or TABLE_CAPTION_PATTERN.match(record["text"])
            or FIGURE_CAPTION_PATTERN.match(record["text"])
        ):
            continue

        total_body_paragraphs += 1

        alignment_name = str(
            paragraph.alignment
        )

        alignments[
            alignment_name
        ] += 1

        spacing = (
            paragraph.paragraph_format.line_spacing
        )

        if isinstance(
            spacing,
            (int, float),
        ):
            spacing_value = round(
                float(spacing),
                2,
            )

            line_spacings[
                str(spacing_value)
            ] += 1

        for run in paragraph.runs:
            if not clean_text(run.text):
                continue

            total_runs += 1

            font_name = get_run_font_name(
                run
            )

            size = safe_pt(
                run.font.size
            )

            if font_name:
                font_names[
                    font_name
                ] += 1

                if (
                    expected_font
                    and font_name.lower()
                    != expected_font
                ):
                    wrong_font_runs += 1

            if size is not None:
                font_sizes[
                    str(size)
                ] += 1

                if abs(
                    size - expected_size
                ) > 0.25:
                    wrong_size_runs += 1

    if total_runs and wrong_font_runs:
        ratio = wrong_font_runs / total_runs

        if ratio >= 0.1:
            add_finding(
                findings,
                "warning",
                "font_mismatch",
                "Jenis font isi belum konsisten",
                (
                    f"{wrong_font_runs} dari {total_runs} run "
                    f"teks menggunakan font berbeda dari "
                    f"{policy.get('font')}."
                ),
                recommendation=(
                    "Samakan font isi dokumen sesuai aturan."
                ),
            )

    if total_runs and wrong_size_runs:
        ratio = wrong_size_runs / total_runs

        if ratio >= 0.1:
            add_finding(
                findings,
                "warning",
                "font_size_mismatch",
                "Ukuran huruf isi belum konsisten",
                (
                    f"{wrong_size_runs} dari {total_runs} run "
                    f"teks tidak menggunakan ukuran "
                    f"{expected_size} pt."
                ),
                recommendation=(
                    "Samakan ukuran huruf isi dokumen."
                ),
            )

    if line_spacings:
        dominant_spacing = float(
            line_spacings.most_common(1)[0][0]
        )

        if abs(
            dominant_spacing - expected_spacing
        ) > 0.15:
            add_finding(
                findings,
                "warning",
                "line_spacing_mismatch",
                "Spasi baris dominan tidak sesuai",
                (
                    f"Spasi dominan {dominant_spacing}, "
                    f"sedangkan aturan {expected_spacing}."
                ),
                recommendation=(
                    "Terapkan spasi baris yang konsisten "
                    "pada paragraf isi."
                ),
            )
    else:
        add_finding(
            findings,
            "info",
            "line_spacing_inherited",
            "Spasi baris menggunakan style bawaan",
            (
                "Sebagian besar paragraf tidak menyimpan nilai "
                "spasi secara langsung dan kemungkinan "
                "mengikuti style Word."
            ),
            recommendation=(
                "Pastikan style Normal memiliki spasi yang "
                "sesuai."
            ),
        )

    expected_margins = {
        "top": float(
            policy.get("margin_top_cm", 3)
        ),
        "bottom": float(
            policy.get("margin_bottom_cm", 3)
        ),
        "left": float(
            policy.get("margin_left_cm", 3)
        ),
        "right": float(
            policy.get("margin_right_cm", 3)
        ),
    }

    section_results: list[dict[str, Any]] = []

    for index, section in enumerate(
        document.sections,
        start=1,
    ):
        actual_margins = {
            "top": safe_cm(
                section.top_margin
            ),
            "bottom": safe_cm(
                section.bottom_margin
            ),
            "left": safe_cm(
                section.left_margin
            ),
            "right": safe_cm(
                section.right_margin
            ),
        }

        section_results.append(
            {
                "section": index,
                "margins_cm": actual_margins,
            }
        )

        for side, expected_value in expected_margins.items():
            actual_value = actual_margins[
                side
            ]

            if actual_value is None:
                continue

            if abs(
                actual_value - expected_value
            ) > 0.15:
                add_finding(
                    findings,
                    "warning",
                    "margin_mismatch",
                    "Margin halaman tidak sesuai",
                    (
                        f"Section {index}, margin {side} "
                        f"{actual_value} cm; aturan "
                        f"{expected_value} cm."
                    ),
                    location=f"Section {index}",
                    recommendation=(
                        "Sesuaikan margin halaman berdasarkan "
                        "template yang digunakan."
                    ),
                )

    return {
        "font_distribution": dict(
            font_names.most_common()
        ),
        "font_size_distribution": dict(
            font_sizes.most_common()
        ),
        "line_spacing_distribution": dict(
            line_spacings.most_common()
        ),
        "alignment_distribution": dict(
            alignments.most_common()
        ),
        "body_paragraphs_checked": (
            total_body_paragraphs
        ),
        "runs_checked": total_runs,
        "wrong_font_runs": wrong_font_runs,
        "wrong_size_runs": wrong_size_runs,
        "sections": section_results,
    }


def calculate_score(
    findings: list[dict[str, Any]],
) -> dict[str, Any]:
    weights = {
        "error": 10,
        "warning": 4,
        "info": 1,
    }

    deductions = sum(
        weights.get(
            finding["severity"],
            1,
        )
        for finding in findings
    )

    score = max(
        0,
        100 - deductions,
    )

    if score >= 90:
        grade = "Sangat Baik"
    elif score >= 80:
        grade = "Baik"
    elif score >= 65:
        grade = "Cukup"
    elif score >= 50:
        grade = "Perlu Perbaikan"
    else:
        grade = "Tidak Memadai"

    severity_counter = Counter(
        finding["severity"]
        for finding in findings
    )

    return {
        "score": score,
        "grade": grade,
        "deductions": deductions,
        "errors": severity_counter.get(
            "error",
            0,
        ),
        "warnings": severity_counter.get(
            "warning",
            0,
        ),
        "information": severity_counter.get(
            "info",
            0,
        ),
    }


def audit_document(
    file_path: str | Path,
    document_type: str = "skripsi",
    custom_policy: dict[str, Any] | None = None,
) -> dict[str, Any]:
    path = Path(file_path)

    normalized_type = document_type.lower().strip()

    if normalized_type not in DEFAULT_POLICIES:
        normalized_type = "skripsi"

    policy = dict(
        DEFAULT_POLICIES[
            normalized_type
        ]
    )

    if custom_policy:
        for key, value in custom_policy.items():
            if value is not None:
                policy[key] = value

    document = Document(path)

    paragraph_records: list[dict[str, Any]] = []
    plain_paragraphs: list[str] = []

    empty_streak = 0
    maximum_empty_streak = 0

    for index, paragraph in enumerate(
        document.paragraphs
    ):
        text = clean_text(
            paragraph.text
        )

        if not text:
            empty_streak += 1
            maximum_empty_streak = max(
                maximum_empty_streak,
                empty_streak,
            )
            continue

        empty_streak = 0
        plain_paragraphs.append(text)

        paragraph_records.append(
            {
                "index": index,
                "text": text,
                "style": (
                    paragraph.style.name
                    if paragraph.style
                    else "Tanpa Style"
                ),
                "paragraph": paragraph,
            }
        )

    findings: list[dict[str, Any]] = []

    front_matter = detect_front_matter(
        plain_paragraphs
    )

    chapters = extract_chapters(
        paragraph_records
    )

    subheadings = extract_subheadings(
        paragraph_records
    )

    table_captions = extract_captions(
        paragraph_records,
        TABLE_CAPTION_PATTERN,
    )

    figure_captions = extract_captions(
        paragraph_records,
        FIGURE_CAPTION_PATTERN,
    )

    appendices = [
        {
            "paragraph_index": record["index"],
            "text": record["text"],
        }
        for record in paragraph_records
        if APPENDIX_PATTERN.match(
            record["text"]
        )
    ]

    checklist: list[dict[str, Any]] = []

    for key in policy.get(
        "required_front_matter",
        [],
    ):
        exists = bool(
            front_matter.get(key)
        )

        checklist.append(
            {
                "item": key.replace(
                    "_",
                    " ",
                ).title(),
                "status": (
                    "passed"
                    if exists
                    else "failed"
                ),
                "message": (
                    "Terdeteksi"
                    if exists
                    else "Belum terdeteksi"
                ),
            }
        )

        if not exists:
            add_finding(
                findings,
                "warning",
                "front_matter_missing",
                "Bagian awal belum lengkap",
                (
                    f"Bagian '{key.replace('_', ' ')}' "
                    "belum terdeteksi."
                ),
                recommendation=(
                    "Tambahkan bagian tersebut sesuai "
                    "pedoman dokumen."
                ),
            )

    checklist.extend(
        audit_chapter_sequence(
            chapters,
            findings,
            list(
                policy.get(
                    "required_chapters",
                    [],
                )
            ),
        )
    )

    audit_subheading_sequence(
        subheadings,
        findings,
    )

    audit_caption_sequence(
        table_captions,
        "tabel",
        findings,
    )

    audit_caption_sequence(
        figure_captions,
        "gambar",
        findings,
    )

    if len(document.tables) > len(
        table_captions
    ):
        difference = (
            len(document.tables)
            - len(table_captions)
        )

        add_finding(
            findings,
            "warning",
            "table_caption_count",
            "Tabel tanpa caption terindikasi",
            (
                f"Dokumen memiliki {len(document.tables)} tabel "
                f"dan {len(table_captions)} caption tabel. "
                f"Selisih: {difference}."
            ),
            recommendation=(
                "Periksa setiap tabel dan tambahkan caption "
                "yang belum tersedia."
            ),
        )

    figure_count = len(
        document.inline_shapes
    )

    if figure_count > len(
        figure_captions
    ):
        difference = (
            figure_count
            - len(figure_captions)
        )

        add_finding(
            findings,
            "warning",
            "figure_caption_count",
            "Gambar tanpa caption terindikasi",
            (
                f"Dokumen memiliki {figure_count} inline image "
                f"dan {len(figure_captions)} caption gambar. "
                f"Selisih: {difference}."
            ),
            recommendation=(
                "Periksa gambar dan tambahkan caption yang "
                "belum tersedia."
            ),
        )

    if maximum_empty_streak >= 4:
        add_finding(
            findings,
            "info",
            "excessive_blank_paragraphs",
            "Paragraf kosong berulang",
            (
                f"Ditemukan hingga {maximum_empty_streak} "
                "paragraf kosong berturut-turut."
            ),
            recommendation=(
                "Gunakan page break atau pengaturan spacing, "
                "bukan paragraf kosong berulang."
            ),
        )

    has_toc_field = contains_field(
        document,
        "TOC",
    )

    has_page_field = contains_field(
        document,
        "PAGE",
    )

    checklist.append(
        {
            "item": "Daftar isi otomatis",
            "status": (
                "passed"
                if has_toc_field
                else "failed"
            ),
            "message": (
                "Field TOC terdeteksi"
                if has_toc_field
                else "Field TOC belum terdeteksi"
            ),
        }
    )

    checklist.append(
        {
            "item": "Nomor halaman otomatis",
            "status": (
                "passed"
                if has_page_field
                else "failed"
            ),
            "message": (
                "Field PAGE terdeteksi"
                if has_page_field
                else "Field PAGE belum terdeteksi"
            ),
        }
    )

    if not has_toc_field:
        add_finding(
            findings,
            "info",
            "toc_field_missing",
            "Daftar isi otomatis belum terdeteksi",
            (
                "Field TOC Microsoft Word belum ditemukan."
            ),
            recommendation=(
                "Gunakan daftar isi otomatis berbasis Heading."
            ),
        )

    if not has_page_field:
        add_finding(
            findings,
            "info",
            "page_field_missing",
            "Nomor halaman otomatis belum terdeteksi",
            (
                "Field PAGE Microsoft Word belum ditemukan."
            ),
            recommendation=(
                "Tambahkan nomor halaman menggunakan field "
                "otomatis Microsoft Word."
            ),
        )

    reference_heading_record = next(
        (
            record
            for record in paragraph_records
            if FRONT_MATTER_PATTERNS[
                "daftar_pustaka"
            ].match(record["text"])
        ),
        None,
    )

    reference_start_index = (
        reference_heading_record["index"]
        if reference_heading_record
        else None
    )

    references = extract_reference_entries(
        paragraph_records
    )

    citations = extract_citations(
        paragraph_records,
        reference_start_index,
    )

    citation_report = (
        audit_citations_and_references(
            citations,
            references,
            findings,
            bool(
                policy.get(
                    "require_references",
                    True,
                )
            ),
        )
    )

    formatting_report = audit_formatting(
        document,
        paragraph_records,
        policy,
        findings,
    )

    score = calculate_score(
        findings
    )

    return {
        "file": {
            "name": path.name,
            "size_bytes": path.stat().st_size,
        },
        "document_type": normalized_type,
        "policy": policy,
        "score": score,
        "summary": {
            "paragraphs_total": len(
                document.paragraphs
            ),
            "paragraphs_non_empty": len(
                paragraph_records
            ),
            "sections_total": len(
                document.sections
            ),
            "tables_total": len(
                document.tables
            ),
            "inline_images_total": figure_count,
            "chapters_total": len(
                chapters
            ),
            "subheadings_total": len(
                subheadings
            ),
            "table_captions_total": len(
                table_captions
            ),
            "figure_captions_total": len(
                figure_captions
            ),
            "appendices_total": len(
                appendices
            ),
            "findings_total": len(
                findings
            ),
        },
        "front_matter": front_matter,
        "structure": {
            "chapters": chapters,
            "subheadings": subheadings,
            "table_captions": table_captions,
            "figure_captions": figure_captions,
            "appendices": appendices,
        },
        "citations": citation_report,
        "formatting": formatting_report,
        "checklist": checklist,
        "findings": findings,
        "limitations": [
            (
                "Pencocokan sitasi dan daftar pustaka "
                "bersifat indikatif."
            ),
            (
                "Audit belum dapat memastikan posisi objek "
                "berdasarkan tampilan halaman Microsoft Word."
            ),
            (
                "Gambar floating atau objek kompleks mungkin "
                "tidak terhitung sebagai inline image."
            ),
            (
                "Kesesuaian substansi ilmiah tidak dinilai "
                "oleh audit format."
            ),
        ],
        "audited_at": utc_now(),
    }