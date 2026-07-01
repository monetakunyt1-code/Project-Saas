from __future__ import annotations

import re
from copy import deepcopy
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph


CHAPTER_PATTERN = re.compile(
    r"^\s*BAB\s+([IVXLCDM]+|\d+)\b",
    re.IGNORECASE,
)

SUBHEADING_PATTERN = re.compile(
    r"^\s*(\d+\.\d+(?:\.\d+)*)\s+(.+)$"
)

TABLE_CAPTION_PATTERN = re.compile(
    r"^\s*Tabel\s+\d+(?:\.\d+)*",
    re.IGNORECASE,
)

FIGURE_CAPTION_PATTERN = re.compile(
    r"^\s*(Gambar|Figure)\s+\d+(?:\.\d+)*",
    re.IGNORECASE,
)

APPENDIX_PATTERN = re.compile(
    r"^\s*Lampiran(?:\s+\d+|\s+[A-Z])?",
    re.IGNORECASE,
)

FRONT_HEADINGS = {
    "HALAMAN JUDUL",
    "LEMBAR PERSETUJUAN",
    "LEMBAR PENGESAHAN",
    "PERNYATAAN KEASLIAN",
    "KATA PENGANTAR",
    "DAFTAR ISI",
    "DAFTAR TABEL",
    "DAFTAR GAMBAR",
    "DAFTAR LAMPIRAN",
    "DAFTAR SINGKATAN",
    "ABSTRAK",
    "ABSTRACT",
    "DAFTAR PUSTAKA",
    "REFERENCES",
}

PRESETS: dict[str, dict[str, Any]] = {
    "skripsi": {
        "font": "Times New Roman",
        "font_size": 12,
        "table_font_size": 10,
        "line_spacing": 2.0,
        "first_line_indent_cm": 1.25,
        "margin_top_cm": 4,
        "margin_bottom_cm": 3,
        "margin_left_cm": 4,
        "margin_right_cm": 3,
        "heading_1_size": 14,
        "heading_2_size": 12,
        "heading_3_size": 12,
    },
    "laporan": {
        "font": "Arial",
        "font_size": 11,
        "table_font_size": 10,
        "line_spacing": 1.5,
        "first_line_indent_cm": 1.25,
        "margin_top_cm": 3,
        "margin_bottom_cm": 3,
        "margin_left_cm": 3,
        "margin_right_cm": 3,
        "heading_1_size": 14,
        "heading_2_size": 12,
        "heading_3_size": 11,
    },
    "jurnal": {
        "font": "Times New Roman",
        "font_size": 11,
        "table_font_size": 9,
        "line_spacing": 1.15,
        "first_line_indent_cm": 0.75,
        "margin_top_cm": 2.5,
        "margin_bottom_cm": 2.5,
        "margin_left_cm": 2.5,
        "margin_right_cm": 2.5,
        "heading_1_size": 13,
        "heading_2_size": 11,
        "heading_3_size": 11,
    },
}


def _set_run_font(
    run: Any,
    font_name: str,
    size_pt: float,
    bold: bool | None = None,
) -> None:
    run.font.name = font_name
    run.font.size = Pt(size_pt)

    if bold is not None:
        run.bold = bold

    run_element = run._element
    run_properties = run_element.get_or_add_rPr()
    run_fonts = run_properties.rFonts

    if run_fonts is None:
        run_fonts = OxmlElement("w:rFonts")
        run_properties.insert(0, run_fonts)

    run_fonts.set(qn("w:ascii"), font_name)
    run_fonts.set(qn("w:hAnsi"), font_name)
    run_fonts.set(qn("w:eastAsia"), font_name)
    run_fonts.set(qn("w:cs"), font_name)


def _set_style_font(
    style: Any,
    font_name: str,
    size_pt: float,
    bold: bool | None = None,
) -> None:
    style.font.name = font_name
    style.font.size = Pt(size_pt)

    if bold is not None:
        style.font.bold = bold

    style_element = style._element
    style_properties = style_element.get_or_add_rPr()
    run_fonts = style_properties.rFonts

    if run_fonts is None:
        run_fonts = OxmlElement("w:rFonts")
        style_properties.insert(0, run_fonts)

    run_fonts.set(qn("w:ascii"), font_name)
    run_fonts.set(qn("w:hAnsi"), font_name)
    run_fonts.set(qn("w:eastAsia"), font_name)
    run_fonts.set(qn("w:cs"), font_name)


def _ensure_custom_styles(
    document: Document,
    rules: dict[str, Any],
) -> None:
    styles = document.styles

    normal = styles["Normal"]
    _set_style_font(
        normal,
        rules["font"],
        rules["font_size"],
        False,
    )

    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = rules["line_spacing"]

    for style_name, size in (
        ("Heading 1", rules["heading_1_size"]),
        ("Heading 2", rules["heading_2_size"]),
        ("Heading 3", rules["heading_3_size"]),
    ):
        style = styles[style_name]
        _set_style_font(
            style,
            rules["font"],
            size,
            True,
        )
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)

    if "Lampiran" not in styles:
        appendix_style = styles.add_style(
            "Lampiran",
            WD_STYLE_TYPE.PARAGRAPH,
        )
    else:
        appendix_style = styles["Lampiran"]

    _set_style_font(
        appendix_style,
        rules["font"],
        rules["heading_2_size"],
        True,
    )


def _apply_page_layout(
    document: Document,
    rules: dict[str, Any],
) -> None:
    for section in document.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(rules["margin_top_cm"])
        section.bottom_margin = Cm(rules["margin_bottom_cm"])
        section.left_margin = Cm(rules["margin_left_cm"])
        section.right_margin = Cm(rules["margin_right_cm"])
        section.header_distance = Cm(1.25)
        section.footer_distance = Cm(1.25)


def _paragraph_level(text: str) -> int | None:
    match = SUBHEADING_PATTERN.match(text)

    if not match:
        return None

    numbering = match.group(1)
    segments = numbering.split(".")

    if len(segments) == 2:
        return 2

    return 3


def _format_paragraphs(
    document: Document,
    rules: dict[str, Any],
) -> dict[str, int]:
    counters = {
        "chapters": 0,
        "subheadings": 0,
        "captions": 0,
        "appendices": 0,
        "body_paragraphs": 0,
    }

    inside_references = False

    for paragraph in document.paragraphs:
        text = re.sub(r"\s+", " ", paragraph.text or "").strip()

        if not text:
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            continue

        upper_text = text.upper()

        if upper_text in {"DAFTAR PUSTAKA", "REFERENCES"}:
            inside_references = True
        elif CHAPTER_PATTERN.match(text):
            inside_references = False

        if CHAPTER_PATTERN.match(text):
            paragraph.style = document.styles["Heading 1"]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.page_break_before = True
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.first_line_indent = Cm(0)
            counters["chapters"] += 1

            for run in paragraph.runs:
                _set_run_font(
                    run,
                    rules["font"],
                    rules["heading_1_size"],
                    True,
                )

            continue

        if upper_text in FRONT_HEADINGS:
            paragraph.style = document.styles["Heading 1"]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.keep_with_next = True

            for run in paragraph.runs:
                _set_run_font(
                    run,
                    rules["font"],
                    rules["heading_1_size"],
                    True,
                )

            continue

        level = _paragraph_level(text)

        if level is not None:
            paragraph.style = document.styles[
                "Heading 2" if level == 2 else "Heading 3"
            ]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.keep_with_next = True
            counters["subheadings"] += 1

            heading_size = (
                rules["heading_2_size"]
                if level == 2
                else rules["heading_3_size"]
            )

            for run in paragraph.runs:
                _set_run_font(
                    run,
                    rules["font"],
                    heading_size,
                    True,
                )

            continue

        if TABLE_CAPTION_PATTERN.match(text):
            try:
                paragraph.style = document.styles["Caption"]
            except KeyError:
                pass

            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.keep_with_next = True
            counters["captions"] += 1

            for run in paragraph.runs:
                _set_run_font(
                    run,
                    rules["font"],
                    rules["table_font_size"],
                    True,
                )

            continue

        if FIGURE_CAPTION_PATTERN.match(text):
            try:
                paragraph.style = document.styles["Caption"]
            except KeyError:
                pass

            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            counters["captions"] += 1

            for run in paragraph.runs:
                _set_run_font(
                    run,
                    rules["font"],
                    rules["table_font_size"],
                    False,
                )

            continue

        if APPENDIX_PATTERN.match(text):
            paragraph.style = document.styles["Lampiran"]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.page_break_before = True
            counters["appendices"] += 1

            for run in paragraph.runs:
                _set_run_font(
                    run,
                    rules["font"],
                    rules["heading_2_size"],
                    True,
                )

            continue

        paragraph.style = document.styles["Normal"]
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = rules["line_spacing"]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        if inside_references:
            paragraph.paragraph_format.left_indent = Cm(1.25)
            paragraph.paragraph_format.first_line_indent = Cm(-1.25)
        else:
            paragraph.paragraph_format.left_indent = Cm(0)
            paragraph.paragraph_format.first_line_indent = Cm(
                rules["first_line_indent_cm"]
            )

        for run in paragraph.runs:
            _set_run_font(
                run,
                rules["font"],
                rules["font_size"],
            )

        counters["body_paragraphs"] += 1

    return counters


def _format_tables(
    document: Document,
    rules: dict[str, Any],
) -> int:
    formatted_cells = 0

    for table in document.tables:
        table.autofit = True

        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                cell.vertical_alignment = (
                    WD_CELL_VERTICAL_ALIGNMENT.CENTER
                )

                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1
                    paragraph.paragraph_format.first_line_indent = Cm(0)

                    if row_index == 0:
                        paragraph.alignment = (
                            WD_ALIGN_PARAGRAPH.CENTER
                        )

                    for run in paragraph.runs:
                        _set_run_font(
                            run,
                            rules["font"],
                            rules["table_font_size"],
                            row_index == 0,
                        )

                formatted_cells += 1

    return formatted_cells


def _add_field(
    paragraph: Paragraph,
    instruction: str,
) -> None:
    run = paragraph.add_run()

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")

    instruction_element = OxmlElement("w:instrText")
    instruction_element.set(
        qn("xml:space"),
        "preserve",
    )
    instruction_element.text = instruction

    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")

    placeholder = OxmlElement("w:t")
    placeholder.text = (
        "Klik kanan lalu pilih Update Field."
    )

    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")

    run._r.append(begin)
    run._r.append(instruction_element)
    run._r.append(separate)
    run._r.append(placeholder)
    run._r.append(end)


def _insert_paragraph_before(
    anchor: Paragraph,
    text: str = "",
) -> Paragraph:
    paragraph_element = OxmlElement("w:p")
    anchor._p.addprevious(paragraph_element)
    paragraph = Paragraph(
        paragraph_element,
        anchor._parent,
    )

    if text:
        paragraph.add_run(text)

    return paragraph


def _find_heading(
    document: Document,
    heading: str,
) -> Paragraph | None:
    target = heading.upper()

    for paragraph in document.paragraphs:
        if paragraph.text.strip().upper() == target:
            return paragraph

    return None


def _field_exists(
    document: Document,
    instruction_fragment: str,
) -> bool:
    instructions = document._element.xpath(".//w:instrText")

    return any(
        instruction_fragment.upper() in (node.text or "").upper()
        for node in instructions
    )


def _insert_field_after(
    paragraph: Paragraph,
    instruction: str,
) -> None:
    new_element = OxmlElement("w:p")
    paragraph._p.addnext(new_element)

    new_paragraph = Paragraph(
        new_element,
        paragraph._parent,
    )

    _add_field(
        new_paragraph,
        instruction,
    )


def _ensure_automatic_lists(
    document: Document,
    rules: dict[str, Any],
    include_toc: bool,
    include_table_list: bool,
    include_figure_list: bool,
    include_appendix_list: bool,
) -> list[str]:
    inserted: list[str] = []

    chapter_anchor = next(
        (
            paragraph
            for paragraph in document.paragraphs
            if CHAPTER_PATTERN.match(paragraph.text.strip())
        ),
        None,
    )

    requests = []

    if include_toc:
        requests.append(
            (
                "DAFTAR ISI",
                'TOC \\o "1-3" \\h \\z \\u',
                "TOC",
            )
        )

    if include_table_list:
        requests.append(
            (
                "DAFTAR TABEL",
                'TOC \\h \\z \\c "Tabel"',
                '\\c "Tabel"',
            )
        )

    if include_figure_list:
        requests.append(
            (
                "DAFTAR GAMBAR",
                'TOC \\h \\z \\c "Gambar"',
                '\\c "Gambar"',
            )
        )

    if include_appendix_list:
        requests.append(
            (
                "DAFTAR LAMPIRAN",
                'TOC \\h \\z \\t "Lampiran,1"',
                '"Lampiran,1"',
            )
        )

    for heading_text, instruction, field_marker in requests:
        heading = _find_heading(
            document,
            heading_text,
        )

        if heading is None and chapter_anchor is not None:
            heading = _insert_paragraph_before(
                chapter_anchor,
                heading_text,
            )
            heading.style = document.styles["Heading 1"]
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

            for run in heading.runs:
                _set_run_font(
                    run,
                    rules["font"],
                    rules["heading_1_size"],
                    True,
                )

        if heading is not None and not _field_exists(
            document,
            field_marker,
        ):
            _insert_field_after(
                heading,
                instruction,
            )
            inserted.append(heading_text)

    return inserted


def _set_page_number_type(
    sect_pr: Any,
    number_format: str,
    start: int = 1,
) -> None:
    page_number_type = sect_pr.find(qn("w:pgNumType"))

    if page_number_type is None:
        page_number_type = OxmlElement("w:pgNumType")
        sect_pr.append(page_number_type)

    page_number_type.set(
        qn("w:fmt"),
        number_format,
    )
    page_number_type.set(
        qn("w:start"),
        str(start),
    )


def _add_page_field_to_footer(
    section: Any,
) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if "PAGE" in "".join(
        node.text or ""
        for node in paragraph._p.xpath(".//w:instrText")
    ).upper():
        return

    _add_field(
        paragraph,
        "PAGE",
    )


def _configure_page_numbers(
    document: Document,
) -> int:
    chapter_paragraph = next(
        (
            paragraph
            for paragraph in document.paragraphs
            if CHAPTER_PATTERN.match(paragraph.text.strip())
        ),
        None,
    )

    if chapter_paragraph is not None:
        previous = chapter_paragraph._p.getprevious()

        if previous is not None:
            previous_properties = previous.get_or_add_pPr()

            if previous_properties.find(qn("w:sectPr")) is None:
                final_section_properties = (
                    document._element.body.sectPr
                )

                if final_section_properties is not None:
                    first_section_properties = deepcopy(
                        final_section_properties
                    )

                    _set_page_number_type(
                        first_section_properties,
                        "lowerRoman",
                        1,
                    )

                    previous_properties.append(
                        first_section_properties
                    )

                    _set_page_number_type(
                        final_section_properties,
                        "decimal",
                        1,
                    )

    for index, section in enumerate(document.sections):
        section.footer.is_linked_to_previous = False
        _add_page_field_to_footer(section)

        section_properties = section._sectPr

        _set_page_number_type(
            section_properties,
            "lowerRoman" if index == 0 and len(document.sections) > 1 else "decimal",
            1,
        )

    return len(document.sections)


def format_document(
    input_path: str | Path,
    output_path: str | Path,
    preset: str = "skripsi",
    custom_rules: dict[str, Any] | None = None,
    include_toc: bool = True,
    include_table_list: bool = True,
    include_figure_list: bool = True,
    include_appendix_list: bool = True,
    include_page_numbers: bool = True,
) -> dict[str, Any]:
    source = Path(input_path)
    destination = Path(output_path)

    selected_preset = preset.lower().strip()

    if selected_preset not in PRESETS:
        selected_preset = "skripsi"

    rules = dict(PRESETS[selected_preset])

    if custom_rules:
        for key, value in custom_rules.items():
            if value is not None and key in rules:
                rules[key] = value

    document = Document(source)

    _ensure_custom_styles(
        document,
        rules,
    )

    _apply_page_layout(
        document,
        rules,
    )

    paragraph_report = _format_paragraphs(
        document,
        rules,
    )

    formatted_cells = _format_tables(
        document,
        rules,
    )

    automatic_lists = _ensure_automatic_lists(
        document,
        rules,
        include_toc,
        include_table_list,
        include_figure_list,
        include_appendix_list,
    )

    section_count = 0

    if include_page_numbers:
        section_count = _configure_page_numbers(
            document
        )

    settings_element = document.settings.element
    update_fields = settings_element.find(qn("w:updateFields"))

    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings_element.append(update_fields)

    update_fields.set(
        qn("w:val"),
        "true",
    )

    destination.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    document.save(destination)

    return {
        "preset": selected_preset,
        "rules": rules,
        "paragraphs": paragraph_report,
        "formatted_table_cells": formatted_cells,
        "automatic_lists_inserted": automatic_lists,
        "page_number_sections": section_count,
        "output_path": str(destination),
    }