from __future__ import annotations

import math
import re
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


CHAPTER_PATTERN = re.compile(
    r"^\s*BAB\s+([IVXLCDM]+|\d+)\b(?:\s*[:.\-]?\s*(.*))?$",
    re.IGNORECASE,
)

ABSTRACT_HEADING_PATTERN = re.compile(
    r"^\s*(ABSTRAK|ABSTRACT)\s*$",
    re.IGNORECASE,
)

REFERENCE_HEADING_PATTERN = re.compile(
    r"^\s*(DAFTAR\s+PUSTAKA|REFERENCES?)\s*$",
    re.IGNORECASE,
)

KEYWORD_PATTERN = re.compile(
    r"^\s*(KATA\s+KUNCI|KEYWORDS?)\s*[:\-]\s*(.+)$",
    re.IGNORECASE,
)

APPENDIX_PATTERN = re.compile(
    r"^\s*LAMPIRAN\b",
    re.IGNORECASE,
)

DIRECT_SECTION_PATTERNS = {
    "introduction": re.compile(
        r"^\s*(PENDAHULUAN|INTRODUCTION)\s*$",
        re.IGNORECASE,
    ),
    "methods": re.compile(
        r"^\s*(METODE(?:\s+PENELITIAN)?|METHODS?)\s*$",
        re.IGNORECASE,
    ),
    "results": re.compile(
        r"^\s*(HASIL(?:\s+DAN\s+PEMBAHASAN)?|"
        r"RESULTS?(?:\s+AND\s+DISCUSSION)?)\s*$",
        re.IGNORECASE,
    ),
    "conclusion": re.compile(
        r"^\s*(KESIMPULAN|CONCLUSION[S]?)\s*$",
        re.IGNORECASE,
    ),
}

ROMAN_VALUES = {
    "I": 1,
    "V": 5,
    "X": 10,
    "L": 50,
    "C": 100,
    "D": 500,
    "M": 1000,
}

CHAPTER_SECTION_MAP = {
    1: "introduction",
    2: "literature",
    3: "methods",
    4: "results",
    5: "conclusion",
}

DEFAULT_RULES = {
    "font": "Times New Roman",
    "font_size": 11.0,
    "heading_1_size": 12.0,
    "title_size": 14.0,
    "line_spacing": 1.15,
    "margin_top_cm": 2.5,
    "margin_bottom_cm": 2.5,
    "margin_left_cm": 2.5,
    "margin_right_cm": 2.5,
}

STOPWORDS = {
    "yang", "dan", "di", "ke", "dari", "untuk", "pada",
    "dengan", "dalam", "adalah", "ini", "itu", "sebagai",
    "oleh", "atau", "serta", "karena", "terhadap", "antara",
    "suatu", "dapat", "akan", "telah", "lebih", "juga",
    "penelitian", "hasil", "berdasarkan", "bahwa", "maka",
    "the", "and", "of", "to", "in", "for", "with", "on",
    "is", "are", "this", "that", "as", "by", "an", "a",
}


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def clean_text(value: str) -> str:
    return re.sub(
        r"\s+",
        " ",
        value or "",
    ).strip()


def word_count(value: str) -> int:
    return len(
        re.findall(
            r"\b[\wÀ-ÖØ-öø-ÿ'-]+\b",
            value or "",
        )
    )


def roman_to_int(value: str) -> int | None:
    normalized = value.upper().strip()

    if normalized.isdigit():
        return int(normalized)

    total = 0
    previous = 0

    for character in reversed(normalized):
        current = ROMAN_VALUES.get(character)

        if current is None:
            return None

        if current < previous:
            total -= current
        else:
            total += current
            previous = current

    return total or None


def split_sentences(text: str) -> list[str]:
    cleaned = clean_text(text)

    if not cleaned:
        return []

    sentences = re.split(
        r"(?<=[.!?])\s+(?=[A-ZÀ-ÖØ-Ý0-9])",
        cleaned,
    )

    return [
        clean_text(sentence)
        for sentence in sentences
        if clean_text(sentence)
    ]


def sentence_tokens(sentence: str) -> list[str]:
    tokens = re.findall(
        r"\b[\wÀ-ÖØ-öø-ÿ'-]{3,}\b",
        sentence.lower(),
    )

    return [
        token
        for token in tokens
        if token not in STOPWORDS
        and not token.isdigit()
    ]


def extractive_summary(
    paragraphs: list[str],
    target_words: int,
) -> list[str]:
    cleaned_paragraphs = [
        clean_text(paragraph)
        for paragraph in paragraphs
        if clean_text(paragraph)
    ]

    if not cleaned_paragraphs:
        return []

    full_text = " ".join(cleaned_paragraphs)

    if target_words <= 0 or word_count(full_text) <= target_words:
        return cleaned_paragraphs

    sentences = split_sentences(full_text)

    if not sentences:
        return cleaned_paragraphs

    frequencies: Counter[str] = Counter()

    for sentence in sentences:
        frequencies.update(
            sentence_tokens(sentence)
        )

    if not frequencies:
        selected: list[str] = []
        current_words = 0

        for sentence in sentences:
            selected.append(sentence)
            current_words += word_count(sentence)

            if current_words >= target_words:
                break

        return [" ".join(selected)]

    maximum_frequency = max(
        frequencies.values()
    )

    normalized_frequencies = {
        word: frequency / maximum_frequency
        for word, frequency in frequencies.items()
    }

    scored_sentences: list[
        tuple[int, float, str, int]
    ] = []

    for index, sentence in enumerate(sentences):
        tokens = sentence_tokens(sentence)

        if not tokens:
            score = 0.0
        else:
            raw_score = sum(
                normalized_frequencies.get(
                    token,
                    0.0,
                )
                for token in tokens
            )

            score = raw_score / math.sqrt(
                max(len(tokens), 1)
            )

        scored_sentences.append(
            (
                index,
                score,
                sentence,
                word_count(sentence),
            )
        )

    ranked = sorted(
        scored_sentences,
        key=lambda item: (
            item[1],
            -item[0],
        ),
        reverse=True,
    )

    selected_indices: set[int] = set()
    selected_words = 0

    for index, _, sentence, sentence_words in ranked:
        if selected_words >= target_words:
            break

        selected_indices.add(index)
        selected_words += sentence_words

    ordered_sentences = [
        sentence
        for index, sentence in enumerate(sentences)
        if index in selected_indices
    ]

    return [" ".join(ordered_sentences)]


def parse_source_document(
    source_path: str | Path,
) -> dict[str, Any]:
    document = Document(source_path)

    sections: dict[str, list[str]] = {
        "front": [],
        "abstract": [],
        "introduction": [],
        "literature": [],
        "methods": [],
        "results": [],
        "conclusion": [],
        "references": [],
    }

    keywords: list[str] = []
    current_section = "front"

    for paragraph in document.paragraphs:
        text = clean_text(
            paragraph.text
        )

        if not text:
            continue

        if APPENDIX_PATTERN.match(text):
            current_section = "appendix"
            continue

        if ABSTRACT_HEADING_PATTERN.match(text):
            current_section = "abstract"
            continue

        if REFERENCE_HEADING_PATTERN.match(text):
            current_section = "references"
            continue

        keyword_match = KEYWORD_PATTERN.match(text)

        if keyword_match:
            keywords = [
                clean_text(item)
                for item in re.split(
                    r"[;,]",
                    keyword_match.group(2),
                )
                if clean_text(item)
            ]
            continue

        direct_section = None

        for section_name, pattern in (
            DIRECT_SECTION_PATTERNS.items()
        ):
            if pattern.match(text):
                direct_section = section_name
                break

        if direct_section:
            current_section = direct_section
            continue

        chapter_match = CHAPTER_PATTERN.match(text)

        if chapter_match:
            chapter_number = roman_to_int(
                chapter_match.group(1)
            )

            current_section = CHAPTER_SECTION_MAP.get(
                chapter_number,
                current_section,
            )
            continue

        if current_section in sections:
            sections[current_section].append(text)

    return {
        "sections": sections,
        "keywords": keywords,
        "tables_total": len(document.tables),
        "images_total": len(document.inline_shapes),
    }


def guess_title(
    front_paragraphs: list[str],
) -> str:
    excluded_terms = (
        "SKRIPSI",
        "TESIS",
        "DISERTASI",
        "UNIVERSITAS",
        "FAKULTAS",
        "PROGRAM STUDI",
        "Diajukan",
        "Disusun",
        "Oleh",
        "NIM",
    )

    candidates: list[str] = []

    for paragraph in front_paragraphs[:40]:
        text = clean_text(paragraph)
        upper_text = text.upper()

        if not 15 <= len(text) <= 300:
            continue

        if any(
            term.upper() in upper_text
            for term in excluded_terms
        ):
            continue

        candidates.append(text)

    if not candidates:
        return "DRAFT ARTIKEL ILMIAH"

    return max(
        candidates,
        key=lambda item: (
            len(item.split()),
            len(item),
        ),
    )


def set_style_font(
    style: Any,
    font_name: str,
    size_pt: float,
    bold: bool | None = None,
) -> None:
    style.font.name = font_name
    style.font.size = Pt(size_pt)

    if bold is not None:
        style.font.bold = bold

    properties = style._element.get_or_add_rPr()
    fonts = properties.rFonts

    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.insert(0, fonts)

    for key in (
        "w:ascii",
        "w:hAnsi",
        "w:eastAsia",
        "w:cs",
    ):
        fonts.set(
            qn(key),
            font_name,
        )


def set_run_font(
    run: Any,
    font_name: str,
    size_pt: float,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = font_name
    run.font.size = Pt(size_pt)

    if bold is not None:
        run.bold = bold

    if italic is not None:
        run.italic = italic

    properties = run._element.get_or_add_rPr()
    fonts = properties.rFonts

    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.insert(0, fonts)

    for key in (
        "w:ascii",
        "w:hAnsi",
        "w:eastAsia",
        "w:cs",
    ):
        fonts.set(
            qn(key),
            font_name,
        )


def configure_document(
    document: Document,
    rules: dict[str, Any],
) -> None:
    for section in document.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(
            float(rules["margin_top_cm"])
        )
        section.bottom_margin = Cm(
            float(rules["margin_bottom_cm"])
        )
        section.left_margin = Cm(
            float(rules["margin_left_cm"])
        )
        section.right_margin = Cm(
            float(rules["margin_right_cm"])
        )

    normal = document.styles["Normal"]

    set_style_font(
        normal,
        str(rules["font"]),
        float(rules["font_size"]),
        False,
    )

    normal.paragraph_format.line_spacing = float(
        rules["line_spacing"]
    )
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    heading_1 = document.styles["Heading 1"]

    set_style_font(
        heading_1,
        str(rules["font"]),
        float(rules["heading_1_size"]),
        True,
    )

    heading_1.paragraph_format.space_before = Pt(6)
    heading_1.paragraph_format.space_after = Pt(6)


def add_body_paragraph(
    document: Document,
    text: str,
    rules: dict[str, Any],
) -> None:
    paragraph = document.add_paragraph()

    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Cm(0.75)
    paragraph.paragraph_format.line_spacing = float(
        rules["line_spacing"]
    )
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(3)

    run = paragraph.add_run(text)

    set_run_font(
        run,
        str(rules["font"]),
        float(rules["font_size"]),
    )


def add_article_section(
    document: Document,
    heading: str,
    paragraphs: list[str],
    rules: dict[str, Any],
) -> None:
    heading_paragraph = document.add_paragraph()
    heading_paragraph.style = document.styles["Heading 1"]
    heading_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    heading_run = heading_paragraph.add_run(heading)

    set_run_font(
        heading_run,
        str(rules["font"]),
        float(rules["heading_1_size"]),
        True,
    )

    if not paragraphs:
        placeholder = document.add_paragraph()
        placeholder.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        run = placeholder.add_run(
            "[Bagian ini belum terdeteksi dari dokumen sumber.]"
        )

        set_run_font(
            run,
            str(rules["font"]),
            float(rules["font_size"]),
            italic=True,
        )

        return

    for text in paragraphs:
        add_body_paragraph(
            document,
            text,
            rules,
        )


def process_section(
    paragraphs: list[str],
    mode: str,
    target_words: int,
) -> list[str]:
    if mode == "extractive":
        return extractive_summary(
            paragraphs,
            target_words,
        )

    return [
        clean_text(paragraph)
        for paragraph in paragraphs
        if clean_text(paragraph)
    ]


def build_journal_article(
    input_path: str | Path,
    output_path: str | Path,
    title: str = "",
    author_name: str = "",
    affiliation: str = "",
    email: str = "",
    keywords: list[str] | None = None,
    reduction_mode: str = "extractive",
    include_literature: bool = True,
    target_abstract_words: int = 200,
    target_introduction_words: int = 1200,
    target_methods_words: int = 700,
    target_results_words: int = 1800,
    target_conclusion_words: int = 300,
    custom_rules: dict[str, Any] | None = None,
) -> dict[str, Any]:
    source_path = Path(input_path)
    destination_path = Path(output_path)

    parsed = parse_source_document(
        source_path
    )

    sections = parsed["sections"]

    selected_title = clean_text(title)

    if not selected_title:
        selected_title = guess_title(
            sections["front"]
        )

    selected_keywords = [
        clean_text(item)
        for item in (keywords or [])
        if clean_text(item)
    ]

    if not selected_keywords:
        selected_keywords = parsed["keywords"]

    introduction_source = list(
        sections["introduction"]
    )

    if include_literature:
        introduction_source.extend(
            sections["literature"]
        )

    targets = {
        "abstract": max(
            int(target_abstract_words),
            50,
        ),
        "introduction": max(
            int(target_introduction_words),
            100,
        ),
        "methods": max(
            int(target_methods_words),
            100,
        ),
        "results": max(
            int(target_results_words),
            100,
        ),
        "conclusion": max(
            int(target_conclusion_words),
            50,
        ),
    }

    normalized_mode = reduction_mode.lower().strip()

    if normalized_mode not in {
        "preserve",
        "extractive",
    }:
        normalized_mode = "extractive"

    processed_sections = {
        "abstract": process_section(
            sections["abstract"],
            normalized_mode,
            targets["abstract"],
        ),
        "introduction": process_section(
            introduction_source,
            normalized_mode,
            targets["introduction"],
        ),
        "methods": process_section(
            sections["methods"],
            normalized_mode,
            targets["methods"],
        ),
        "results": process_section(
            sections["results"],
            normalized_mode,
            targets["results"],
        ),
        "conclusion": process_section(
            sections["conclusion"],
            normalized_mode,
            targets["conclusion"],
        ),
        "references": [
            clean_text(paragraph)
            for paragraph in sections["references"]
            if clean_text(paragraph)
        ],
    }

    rules = dict(DEFAULT_RULES)

    if custom_rules:
        for key, value in custom_rules.items():
            if (
                value is not None
                and key in rules
            ):
                rules[key] = value

    document = Document()

    configure_document(
        document,
        rules,
    )

    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    title_run = title_paragraph.add_run(
        selected_title.upper()
    )

    set_run_font(
        title_run,
        str(rules["font"]),
        float(rules["title_size"]),
        True,
    )

    author_paragraph = document.add_paragraph()
    author_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    author_lines = [
        clean_text(author_name) or "Nama Penulis",
        clean_text(affiliation) or "Afiliasi",
        clean_text(email) or "Email",
    ]

    author_run = author_paragraph.add_run(
        "\n".join(author_lines)
    )

    set_run_font(
        author_run,
        str(rules["font"]),
        float(rules["font_size"]),
    )

    add_article_section(
        document,
        "Abstrak",
        processed_sections["abstract"],
        rules,
    )

    keyword_paragraph = document.add_paragraph()
    keyword_paragraph.paragraph_format.space_after = Pt(6)

    keyword_label = keyword_paragraph.add_run(
        "Kata kunci: "
    )

    set_run_font(
        keyword_label,
        str(rules["font"]),
        float(rules["font_size"]),
        True,
    )

    keyword_value = keyword_paragraph.add_run(
        "; ".join(selected_keywords)
        if selected_keywords
        else "[Masukkan 3–5 kata kunci]"
    )

    set_run_font(
        keyword_value,
        str(rules["font"]),
        float(rules["font_size"]),
    )

    add_article_section(
        document,
        "Pendahuluan",
        processed_sections["introduction"],
        rules,
    )

    add_article_section(
        document,
        "Metode Penelitian",
        processed_sections["methods"],
        rules,
    )

    add_article_section(
        document,
        "Hasil dan Pembahasan",
        processed_sections["results"],
        rules,
    )

    add_article_section(
        document,
        "Kesimpulan",
        processed_sections["conclusion"],
        rules,
    )

    heading = document.add_paragraph()
    heading.style = document.styles["Heading 1"]

    heading_run = heading.add_run(
        "Daftar Pustaka"
    )

    set_run_font(
        heading_run,
        str(rules["font"]),
        float(rules["heading_1_size"]),
        True,
    )

    if processed_sections["references"]:
        for reference in processed_sections["references"]:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(1.0)
            paragraph.paragraph_format.first_line_indent = Cm(-1.0)
            paragraph.paragraph_format.space_after = Pt(3)

            run = paragraph.add_run(reference)

            set_run_font(
                run,
                str(rules["font"]),
                float(rules["font_size"]),
            )
    else:
        paragraph = document.add_paragraph()

        run = paragraph.add_run(
            "[Daftar pustaka belum terdeteksi.]"
        )

        set_run_font(
            run,
            str(rules["font"]),
            float(rules["font_size"]),
            italic=True,
        )

    destination_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    document.save(
        destination_path
    )

    source_counts = {
        key: word_count(
            " ".join(value)
        )
        for key, value in {
            "abstract": sections["abstract"],
            "introduction": introduction_source,
            "methods": sections["methods"],
            "results": sections["results"],
            "conclusion": sections["conclusion"],
            "references": sections["references"],
        }.items()
    }

    output_counts = {
        key: word_count(
            " ".join(value)
        )
        for key, value in processed_sections.items()
    }

    warnings: list[str] = []

    for section_name in (
        "abstract",
        "introduction",
        "methods",
        "results",
        "conclusion",
    ):
        if source_counts[section_name] == 0:
            warnings.append(
                f"Bagian {section_name} belum terdeteksi."
            )

    if not processed_sections["references"]:
        warnings.append(
            "Daftar pustaka belum terdeteksi."
        )

    if parsed["tables_total"]:
        warnings.append(
            (
                f"Dokumen sumber memiliki "
                f"{parsed['tables_total']} tabel. "
                "Tabel belum dipindahkan otomatis ke draft jurnal."
            )
        )

    if parsed["images_total"]:
        warnings.append(
            (
                f"Dokumen sumber memiliki "
                f"{parsed['images_total']} gambar. "
                "Gambar belum dipindahkan otomatis ke draft jurnal."
            )
        )

    return {
        "title": selected_title,
        "author": {
            "name": clean_text(author_name),
            "affiliation": clean_text(affiliation),
            "email": clean_text(email),
        },
        "keywords": selected_keywords,
        "reduction_mode": normalized_mode,
        "include_literature": include_literature,
        "targets": targets,
        "source_word_counts": source_counts,
        "output_word_counts": output_counts,
        "source_total_words": sum(
            source_counts.values()
        ),
        "output_total_words": sum(
            output_counts.values()
        ),
        "rules": rules,
        "warnings": warnings,
        "output_path": str(
            destination_path
        ),
        "created_at": utc_now(),
    }