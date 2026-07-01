from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


CHAPTER_PATTERN = re.compile(
    r"^\s*BAB\s+([IVXLCDM]+|\d+)\b",
    re.IGNORECASE,
)

ABSTRACT_PATTERN = re.compile(
    r"^\s*(ABSTRAK|ABSTRACT)\s*$",
    re.IGNORECASE,
)

REFERENCE_PATTERN = re.compile(
    r"^\s*(DAFTAR PUSTAKA|REFERENCES?)\s*$",
    re.IGNORECASE,
)

ROMAN_TO_SECTION = {
    "I": "pendahuluan",
    "1": "pendahuluan",
    "II": "tinjauan_pustaka",
    "2": "tinjauan_pustaka",
    "III": "metode",
    "3": "metode",
    "IV": "hasil_pembahasan",
    "4": "hasil_pembahasan",
    "V": "kesimpulan",
    "5": "kesimpulan",
}


def _clean_text(value: str) -> str:
    return re.sub(
        r"\s+",
        " ",
        value or "",
    ).strip()


def _collect_sections(
    document: Document,
) -> dict[str, list[str]]:
    result: dict[str, list[str]] = {
        "front": [],
        "abstrak": [],
        "pendahuluan": [],
        "tinjauan_pustaka": [],
        "metode": [],
        "hasil_pembahasan": [],
        "kesimpulan": [],
        "referensi": [],
    }

    current_section = "front"

    for paragraph in document.paragraphs:
        text = _clean_text(paragraph.text)

        if not text:
            continue

        if ABSTRACT_PATTERN.match(text):
            current_section = "abstrak"
            continue

        if REFERENCE_PATTERN.match(text):
            current_section = "referensi"
            continue

        chapter_match = CHAPTER_PATTERN.match(text)

        if chapter_match:
            chapter_number = chapter_match.group(1).upper()
            current_section = ROMAN_TO_SECTION.get(
                chapter_number,
                current_section,
            )
            continue

        result.setdefault(
            current_section,
            [],
        ).append(text)

    return result


def _guess_title(
    sections: dict[str, list[str]],
) -> str:
    candidates = sections.get("front", [])

    for candidate in candidates[:20]:
        upper = candidate.upper()

        if (
            len(candidate) >= 15
            and len(candidate) <= 250
            and "UNIVERSITAS" not in upper
            and "FAKULTAS" not in upper
            and "PROGRAM STUDI" not in upper
            and "SKRIPSI" not in upper
        ):
            return candidate

    return "DRAFT ARTIKEL ILMIAH"


def _add_section(
    document: Document,
    heading: str,
    paragraphs: list[str],
) -> int:
    document.add_heading(
        heading,
        level=1,
    )

    total = 0

    if not paragraphs:
        paragraph = document.add_paragraph()
        placeholder_run = paragraph.add_run(
            "[Bagian ini belum terdeteksi dari dokumen sumber.]"
        )
        placeholder_run.italic = True
        return total

    for text in paragraphs:
        paragraph = document.add_paragraph(text)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.first_line_indent = Cm(0.75)
        paragraph.paragraph_format.line_spacing = 1.15
        total += 1

    return total


def convert_to_journal(
    input_path: str | Path,
    output_path: str | Path,
) -> dict[str, Any]:
    source = Path(input_path)
    destination = Path(output_path)

    source_document = Document(source)
    sections = _collect_sections(
        source_document
    )

    output_document = Document()

    for section in output_document.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    normal_style = output_document.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(11)

    title = _guess_title(sections)

    title_paragraph = output_document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    title_run = title_paragraph.add_run(
        title.upper()
    )
    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(14)

    author_paragraph = output_document.add_paragraph(
        "Nama Penulis\nAfiliasi\nEmail"
    )
    author_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    counts: dict[str, int] = {}

    counts["abstrak"] = _add_section(
        output_document,
        "Abstrak",
        sections["abstrak"],
    )

    keywords = output_document.add_paragraph(
        "Kata kunci: kata kunci 1; kata kunci 2; kata kunci 3"
    )
    keywords.runs[0].bold = True

    counts["pendahuluan"] = _add_section(
        output_document,
        "Pendahuluan",
        sections["pendahuluan"],
    )

    counts["metode"] = _add_section(
        output_document,
        "Metode Penelitian",
        sections["metode"],
    )

    counts["hasil_pembahasan"] = _add_section(
        output_document,
        "Hasil dan Pembahasan",
        sections["hasil_pembahasan"],
    )

    counts["kesimpulan"] = _add_section(
        output_document,
        "Kesimpulan",
        sections["kesimpulan"],
    )

    counts["referensi"] = _add_section(
        output_document,
        "Daftar Pustaka",
        sections["referensi"],
    )

    destination.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    output_document.save(destination)

    return {
        "title": title,
        "sections": counts,
        "output_path": str(destination),
        "warning": (
            "Draft jurnal dibuat melalui pemetaan struktur. "
            "Isi belum diringkas menggunakan AI."
        ),
    }