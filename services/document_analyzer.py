from __future__ import annotations

import re
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document


CHAPTER_PATTERN = re.compile(
    r"^\s*BAB\s+([IVXLCDM]+|\d+)\b",
    re.IGNORECASE,
)

SUBHEADING_PATTERN = re.compile(
    r"^\s*(\d+\.\d+(?:\.\d+)*)\s+(.+)$"
)

TABLE_CAPTION_PATTERN = re.compile(
    r"^\s*Tabel\s+\d+(?:\.\d+)*",
    re.IGNORECASE,
)

FIGURE_CAPTION_PATTERN = re.compile(
    r"^\s*(Gambar|Figure)\s+\d+(?:\.\d+)*",
    re.IGNORECASE,
)

APPENDIX_PATTERN = re.compile(
    r"^\s*Lampiran(?:\s+\d+|\s+[A-Z])?",
    re.IGNORECASE,
)

REFERENCE_PATTERN = re.compile(
    r"^\s*(DAFTAR\s+PUSTAKA|REFERENCES?)\s*$",
    re.IGNORECASE,
)

ABSTRACT_PATTERN = re.compile(
    r"^\s*(ABSTRAK|ABSTRACT)\s*$",
    re.IGNORECASE,
)

FRONT_MATTER_HEADINGS = {
    "HALAMAN JUDUL",
    "LEMBAR PERSETUJUAN",
    "LEMBAR PENGESAHAN",
    "PERNYATAAN KEASLIAN",
    "KATA PENGANTAR",
    "DAFTAR ISI",
    "DAFTAR TABEL",
    "DAFTAR GAMBAR",
    "DAFTAR LAMPIRAN",
    "DAFTAR SINGKATAN",
    "ABSTRAK",
    "ABSTRACT",
}


def _clean_text(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def _font_name(run: Any) -> str | None:
    if run.font.name:
        return run.font.name

    try:
        return run._element.rPr.rFonts.get(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii"
        )
    except (AttributeError, TypeError):
        return None


def _font_size(run: Any) -> float | None:
    if run.font.size is None:
        return None

    return round(run.font.size.pt, 2)


def _contains_field(document: Document, field_name: str) -> bool:
    instructions = document._element.xpath(".//w:instrText")

    return any(
        field_name.upper() in (node.text or "").upper()
        for node in instructions
    )


def analyze_document(file_path: str | Path) -> dict[str, Any]:
    path = Path(file_path)
    document = Document(path)

    chapters: list[dict[str, Any]] = []
    subheadings: list[dict[str, Any]] = []
    table_captions: list[str] = []
    figure_captions: list[str] = []
    appendices: list[str] = []
    front_matter: list[str] = []

    font_names: Counter[str] = Counter()
    font_sizes: Counter[str] = Counter()
    styles: Counter[str] = Counter()

    non_empty_paragraphs = 0
    empty_streak = 0
    maximum_empty_streak = 0

    issues: list[dict[str, str]] = []

    for index, paragraph in enumerate(document.paragraphs):
        text = _clean_text(paragraph.text)
        style_name = paragraph.style.name if paragraph.style else "Tanpa Style"
        styles[style_name] += 1

        if not text:
            empty_streak += 1
            maximum_empty_streak = max(
                maximum_empty_streak,
                empty_streak,
            )
            continue

        empty_streak = 0
        non_empty_paragraphs += 1

        upper_text = text.upper()

        chapter_match = CHAPTER_PATTERN.match(text)

        if chapter_match:
            chapters.append(
                {
                    "paragraph_index": index,
                    "number": chapter_match.group(1),
                    "text": text,
                    "style": style_name,
                }
            )

        subheading_match = SUBHEADING_PATTERN.match(text)

        if subheading_match:
            subheadings.append(
                {
                    "paragraph_index": index,
                    "number": subheading_match.group(1),
                    "text": text,
                    "style": style_name,
                }
            )

        if TABLE_CAPTION_PATTERN.match(text):
            table_captions.append(text)

        if FIGURE_CAPTION_PATTERN.match(text):
            figure_captions.append(text)

        if APPENDIX_PATTERN.match(text):
            appendices.append(text)

        if upper_text in FRONT_MATTER_HEADINGS:
            front_matter.append(text)

        for run in paragraph.runs:
            if not _clean_text(run.text):
                continue

            name = _font_name(run)

            if name:
                font_names[name] += 1

            size = _font_size(run)

            if size is not None:
                font_sizes[str(size)] += 1

    section_data: list[dict[str, Any]] = []

    for index, section in enumerate(document.sections):
        section_data.append(
            {
                "section": index + 1,
                "top_margin_cm": round(section.top_margin.cm, 2),
                "bottom_margin_cm": round(section.bottom_margin.cm, 2),
                "left_margin_cm": round(section.left_margin.cm, 2),
                "right_margin_cm": round(section.right_margin.cm, 2),
                "page_width_cm": round(section.page_width.cm, 2),
                "page_height_cm": round(section.page_height.cm, 2),
            }
        )

    if not chapters:
        issues.append(
            {
                "severity": "warning",
                "code": "chapter_not_found",
                "message": "Judul BAB belum terdeteksi.",
            }
        )

    if maximum_empty_streak >= 4:
        issues.append(
            {
                "severity": "warning",
                "code": "excessive_blank_paragraphs",
                "message": (
                    f"Ditemukan hingga {maximum_empty_streak} "
                    "paragraf kosong berturut-turut."
                ),
            }
        )

    if len(font_names) > 3:
        issues.append(
            {
                "severity": "warning",
                "code": "font_inconsistency",
                "message": (
                    "Dokumen menggunakan lebih dari tiga jenis font."
                ),
            }
        )

    if document.tables and not table_captions:
        issues.append(
            {
                "severity": "warning",
                "code": "table_caption_missing",
                "message": (
                    "Dokumen memiliki tabel, tetapi caption tabel "
                    "tidak terdeteksi."
                ),
            }
        )

    if not _contains_field(document, "TOC"):
        issues.append(
            {
                "severity": "info",
                "code": "toc_field_missing",
                "message": (
                    "Field daftar isi otomatis Microsoft Word "
                    "belum terdeteksi."
                ),
            }
        )

    if not _contains_field(document, "PAGE"):
        issues.append(
            {
                "severity": "info",
                "code": "page_field_missing",
                "message": (
                    "Field nomor halaman otomatis belum terdeteksi."
                ),
            }
        )

    return {
        "file": {
            "name": path.name,
            "size_bytes": path.stat().st_size,
        },
        "summary": {
            "paragraphs_total": len(document.paragraphs),
            "paragraphs_non_empty": non_empty_paragraphs,
            "tables_total": len(document.tables),
            "sections_total": len(document.sections),
            "chapters_total": len(chapters),
            "subheadings_total": len(subheadings),
            "table_captions_total": len(table_captions),
            "figure_captions_total": len(figure_captions),
            "appendices_total": len(appendices),
            "issues_total": len(issues),
        },
        "structure": {
            "chapters": chapters,
            "subheadings": subheadings,
            "front_matter": front_matter,
            "table_captions": table_captions,
            "figure_captions": figure_captions,
            "appendices": appendices,
        },
        "formatting": {
            "font_names": dict(font_names.most_common()),
            "font_sizes_pt": dict(font_sizes.most_common()),
            "paragraph_styles": dict(styles.most_common()),
            "sections": section_data,
            "has_toc_field": _contains_field(document, "TOC"),
            "has_page_field": _contains_field(document, "PAGE"),
        },
        "issues": issues,
        "analyzed_at": datetime.now(timezone.utc).isoformat(),
    }