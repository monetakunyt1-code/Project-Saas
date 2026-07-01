from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


BASE_DIR = Path(__file__).resolve().parent.parent

ACCEPTANCE_DIR = (
    BASE_DIR
    / "storage"
    / "acceptance"
)

FIXTURE_DIR = (
    ACCEPTANCE_DIR
    / "fixtures"
)

OUTPUT_DIR = (
    ACCEPTANCE_DIR
    / "outputs"
)

FIXTURE_DIR.mkdir(
    parents=True,
    exist_ok=True,
)

OUTPUT_DIR.mkdir(
    parents=True,
    exist_ok=True,
)


def configure_document(
    document: Document,
) -> None:
    section = document.sections[0]

    section.top_margin = Cm(3)
    section.bottom_margin = Cm(3)
    section.left_margin = Cm(4)
    section.right_margin = Cm(3)

    normal_style = document.styles["Normal"]

    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(12)

    paragraph_format = (
        normal_style.paragraph_format
    )

    paragraph_format.line_spacing = 1.5
    paragraph_format.first_line_indent = Cm(1.25)


def create_simple_document(
    path: Path,
) -> None:
    document = Document()

    configure_document(document)

    title = document.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = title.add_run(
        "DOKUMEN UJI SEDERHANA"
    )

    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)

    document.add_paragraph(
        "Dokumen ini digunakan untuk menguji fungsi "
        "pemformatan dasar pada aplikasi DocuRapi."
    )

    document.add_heading(
        "Pendahuluan",
        level=1,
    )

    document.add_paragraph(
        "Teknologi pengolahan dokumen dapat membantu "
        "pengguna menerapkan format yang konsisten."
    )

    document.add_heading(
        "Kesimpulan",
        level=1,
    )

    document.add_paragraph(
        "Pengujian sederhana selesai dilakukan."
    )

    document.save(path)


def create_academic_document(
    path: Path,
) -> None:
    document = Document()

    configure_document(document)

    title = document.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = title.add_run(
        "PENGEMBANGAN SISTEM PEMFORMATAN "
        "DOKUMEN AKADEMIK"
    )

    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)

    document.add_paragraph(
        "Nama Penulis"
    ).alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    document.add_heading(
        "Abstrak",
        level=1,
    )

    document.add_paragraph(
        "Penelitian ini bertujuan mengembangkan sistem "
        "pemformatan dokumen akademik. Metode penelitian "
        "menggunakan pendekatan pengembangan perangkat "
        "lunak. Hasil menunjukkan bahwa otomatisasi dapat "
        "meningkatkan konsistensi format dokumen."
    )

    keywords = document.add_paragraph()

    keywords.add_run(
        "Kata kunci: "
    ).bold = True

    keywords.add_run(
        "dokumen akademik, otomatisasi, pemformatan"
    )

    document.add_heading(
        "BAB I PENDAHULUAN",
        level=1,
    )

    document.add_heading(
        "1.1 Latar Belakang",
        level=2,
    )

    document.add_paragraph(
        "Konsistensi format merupakan bagian penting "
        "dari penyusunan karya ilmiah (Sari, 2024). "
        "Kesalahan format dapat mengurangi keterbacaan "
        "dan kualitas penyajian dokumen."
    )

    document.add_heading(
        "1.2 Rumusan Masalah",
        level=2,
    )

    document.add_paragraph(
        "Bagaimana sistem otomatis dapat membantu "
        "pemformatan dokumen akademik?"
    )

    document.add_heading(
        "BAB II TINJAUAN PUSTAKA",
        level=1,
    )

    document.add_paragraph(
        "Pemrosesan dokumen merupakan aktivitas untuk "
        "mengatur struktur, gaya, dan tampilan teks "
        "(Putra & Lestari, 2023)."
    )

    document.add_heading(
        "BAB III METODE PENELITIAN",
        level=1,
    )

    document.add_paragraph(
        "Penelitian menggunakan metode pengembangan "
        "dengan tahapan analisis, implementasi, dan uji."
    )

    table = document.add_table(
        rows=1,
        cols=3,
    )

    table.style = "Table Grid"

    headers = table.rows[0].cells

    headers[0].text = "No."
    headers[1].text = "Tahapan"
    headers[2].text = "Hasil"

    rows = [
        ("1", "Analisis", "Kebutuhan sistem"),
        ("2", "Implementasi", "Aplikasi"),
        ("3", "Pengujian", "Laporan uji"),
    ]

    for number, stage, result in rows:
        cells = table.add_row().cells
        cells[0].text = number
        cells[1].text = stage
        cells[2].text = result

    document.add_paragraph(
        "Tabel 1. Tahapan pengembangan sistem"
    )

    document.add_heading(
        "BAB IV HASIL DAN PEMBAHASAN",
        level=1,
    )

    document.add_paragraph(
        "Sistem berhasil mengidentifikasi struktur "
        "dokumen dan menerapkan konfigurasi format."
    )

    document.add_heading(
        "BAB V KESIMPULAN",
        level=1,
    )

    document.add_paragraph(
        "Otomatisasi pemformatan berpotensi membantu "
        "penyusunan dokumen akademik secara konsisten."
    )

    document.add_heading(
        "DAFTAR PUSTAKA",
        level=1,
    )

    document.add_paragraph(
        "Putra, A., & Lestari, D. (2023). "
        "Pemrosesan Dokumen Digital. Jakarta: "
        "Penerbit Akademik."
    )

    document.add_paragraph(
        "Sari, N. (2024). Otomatisasi Dokumen "
        "Akademik. Bandung: Media Ilmiah."
    )

    document.save(path)


def create_inconsistent_document(
    path: Path,
) -> None:
    document = Document()

    section = document.sections[0]

    section.top_margin = Cm(1)
    section.bottom_margin = Cm(4.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(2)

    paragraph = document.add_paragraph()

    first_run = paragraph.add_run(
        "DOKUMEN DENGAN FORMAT "
    )

    first_run.font.name = "Arial"
    first_run.font.size = Pt(18)
    first_run.bold = True

    second_run = paragraph.add_run(
        "TIDAK KONSISTEN"
    )

    second_run.font.name = "Calibri"
    second_run.font.size = Pt(10)

    document.add_heading(
        "BAB I PENDAHULUAN",
        level=1,
    )

    paragraph = document.add_paragraph(
        "Paragraf pertama menggunakan font dan "
        "ukuran yang tidak konsisten."
    )

    paragraph.paragraph_format.line_spacing = 1

    paragraph = document.add_paragraph(
        "Paragraf kedua memiliki spasi yang berbeda "
        "dan menyebut sumber yang tidak ada pada daftar "
        "pustaka (TidakAda, 2099)."
    )

    paragraph.paragraph_format.line_spacing = 2.5
    paragraph.paragraph_format.left_indent = Cm(2)

    document.add_heading(
        "BAB I PENDAHULUAN",
        level=1,
    )

    document.add_heading(
        "1.3 Subbab Melompat",
        level=2,
    )

    document.add_paragraph(
        "Penomoran subbab sengaja dibuat tidak berurutan."
    )

    document.add_paragraph(
        "Tabel 3. Judul tabel tanpa objek tabel."
    )

    document.add_heading(
        "KESIMPULAN",
        level=1,
    )

    document.add_paragraph(
        "Dokumen ini sengaja mengandung sejumlah masalah."
    )

    document.save(path)


def create_batch_documents() -> list[Path]:
    paths: list[Path] = []

    for index in range(1, 4):
        path = (
            FIXTURE_DIR
            / f"batch_test_{index:02d}.docx"
        )

        document = Document()

        configure_document(document)

        document.add_heading(
            f"Dokumen Batch {index}",
            level=1,
        )

        document.add_paragraph(
            "Dokumen ini merupakan bagian dari "
            f"pengujian batch nomor {index}."
        )

        document.add_heading(
            "Isi Dokumen",
            level=2,
        )

        document.add_paragraph(
            "Konten dokumen batch dibuat secara otomatis "
            "oleh Acceptance Factory DocuRapi."
        )

        document.save(path)

        paths.append(path)

    return paths


def create_acceptance_fixtures() -> dict[str, Any]:
    simple_path = (
        FIXTURE_DIR
        / "simple_document.docx"
    )

    academic_path = (
        FIXTURE_DIR
        / "academic_document.docx"
    )

    inconsistent_path = (
        FIXTURE_DIR
        / "inconsistent_document.docx"
    )

    create_simple_document(
        simple_path
    )

    create_academic_document(
        academic_path
    )

    create_inconsistent_document(
        inconsistent_path
    )

    batch_paths = create_batch_documents()

    fixtures = [
        simple_path,
        academic_path,
        inconsistent_path,
        *batch_paths,
    ]

    return {
        "fixture_directory": str(
            FIXTURE_DIR
        ),
        "output_directory": str(
            OUTPUT_DIR
        ),
        "fixtures": [
            {
                "name": path.name,
                "path": str(path),
                "size_bytes": (
                    path.stat().st_size
                ),
            }
            for path in fixtures
        ],
        "simple_document": str(
            simple_path
        ),
        "academic_document": str(
            academic_path
        ),
        "inconsistent_document": str(
            inconsistent_path
        ),
        "batch_documents": [
            str(path)
            for path in batch_paths
        ],
    }
