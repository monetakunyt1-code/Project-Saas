from pathlib import Path
from docx import Document
from formatter import format_document

root = Path(__file__).resolve().parent
input_file = root / "storage" / "smoke_input.docx"
output_file = root / "storage" / "smoke_output.docx"
doc = Document()
doc.add_paragraph("JUDUL DOKUMEN")
doc.add_paragraph("ABSTRAK")
doc.add_paragraph("Contoh paragraf abstrak untuk pengujian aplikasi.")
doc.add_paragraph("BAB I PENDAHULUAN")
doc.add_paragraph("1.1 Latar Belakang")
doc.add_paragraph("Contoh isi paragraf utama untuk memastikan format berjalan.")
doc.add_paragraph("Tabel 1.1 Contoh Tabel")
table = doc.add_table(rows=2, cols=2)
table.cell(0,0).text = "Kolom A"
table.cell(0,1).text = "Kolom B"
table.cell(1,0).text = "1"
table.cell(1,1).text = "2"
doc.add_paragraph("DAFTAR PUSTAKA")
doc.add_paragraph("Contoh, A. (2026). Judul referensi.")
doc.save(input_file)
report = format_document(input_file, output_file, "skripsi", True)
assert output_file.exists() and output_file.stat().st_size > 0
assert report["paragraphs_processed"] > 0
input_file.unlink(missing_ok=True)
output_file.unlink(missing_ok=True)
print("SMOKE TEST BERHASIL")
