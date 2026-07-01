from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

from app import app
from config import DATABASE_PATH
from database import initialize_database
from services.document_analyzer import analyze_document
from services.final_formatter import format_document
from services.journal_converter import convert_to_journal
from services.template_manager import extract_template_rules


root = Path(__file__).resolve().parent
storage = root / "storage"
storage.mkdir(parents=True, exist_ok=True)

initialize_database()

if not DATABASE_PATH.exists():
    raise RuntimeError("Database SQLite gagal dibuat.")

required_routes = {
    "/",
    "/api/health",
    "/api/process",
    "/api/jobs",
    "/api/templates",
}

available_routes = {
    route.path
    for route in app.routes
}

missing_routes = required_routes - available_routes

if missing_routes:
    raise RuntimeError(
        "Endpoint belum lengkap: "
        + ", ".join(sorted(missing_routes))
    )


with TemporaryDirectory(
    prefix="docurapi_check_",
    dir=storage,
) as temp_directory:
    temp = Path(temp_directory)

    source = temp / "contoh.docx"
    formatted = temp / "contoh_rapi.docx"
    journal = temp / "contoh_jurnal.docx"

    document = Document()

    document.add_paragraph("HALAMAN JUDUL")
    document.add_paragraph(
        "PENGEMBANGAN SISTEM PENGELOLAAN DOKUMEN AKADEMIK"
    )

    document.add_paragraph("ABSTRAK")
    document.add_paragraph(
        "Penelitian ini membahas pengelolaan dokumen "
        "akademik secara otomatis."
    )

    document.add_paragraph("DAFTAR ISI")

    document.add_paragraph("BAB I PENDAHULUAN")
    document.add_paragraph("1.1 Latar Belakang")
    document.add_paragraph(
        "Dokumen akademik membutuhkan format yang "
        "sistematis dan konsisten."
    )

    document.add_paragraph("Tabel 1.1 Hasil Pemeriksaan")

    table = document.add_table(
        rows=2,
        cols=2,
    )

    table.cell(0, 0).text = "Komponen"
    table.cell(0, 1).text = "Status"
    table.cell(1, 0).text = "Formatter"
    table.cell(1, 1).text = "Aktif"

    document.add_paragraph("BAB II TINJAUAN PUSTAKA")
    document.add_paragraph("2.1 Pengelolaan Dokumen")
    document.add_paragraph(
        "Pengelolaan dokumen dilakukan berdasarkan aturan."
    )

    document.add_paragraph("BAB III METODE PENELITIAN")
    document.add_paragraph("3.1 Metode")
    document.add_paragraph(
        "Metode yang digunakan adalah pengembangan sistem."
    )

    document.add_paragraph("BAB IV HASIL DAN PEMBAHASAN")
    document.add_paragraph("4.1 Hasil")
    document.add_paragraph(
        "Sistem berhasil mengenali struktur dokumen."
    )

    document.add_paragraph("BAB V KESIMPULAN")
    document.add_paragraph("5.1 Kesimpulan")
    document.add_paragraph(
        "DocuRapi dapat digunakan untuk memproses dokumen."
    )

    document.add_paragraph("DAFTAR PUSTAKA")
    document.add_paragraph(
        "Penulis, A. (2026). Pengelolaan Dokumen Akademik."
    )

    document.add_paragraph("Lampiran 1 Hasil Pengujian")

    document.save(source)

    before = analyze_document(source)

    format_result = format_document(
        input_path=source,
        output_path=formatted,
        preset="skripsi",
        include_toc=True,
        include_table_list=True,
        include_figure_list=True,
        include_appendix_list=True,
        include_page_numbers=True,
    )

    if not formatted.exists():
        raise RuntimeError(
            "Formatter tidak menghasilkan dokumen."
        )

    Document(formatted)

    after = analyze_document(formatted)

    journal_result = convert_to_journal(
        input_path=source,
        output_path=journal,
    )

    if not journal.exists():
        raise RuntimeError(
            "Konverter jurnal tidak menghasilkan dokumen."
        )

    Document(journal)

    template_rules = extract_template_rules(source)

    if not template_rules.get("font"):
        raise RuntimeError(
            "Pembacaan template tidak berhasil."
        )

    print("")
    print("============================================")
    print("FINAL CHECK PASSED")
    print("============================================")
    print("Paragraf sumber :", before["summary"]["paragraphs_total"])
    print("Paragraf hasil  :", after["summary"]["paragraphs_total"])
    print("Tabel hasil     :", after["summary"]["tables_total"])
    print("Preset          :", format_result["preset"])
    print("Judul jurnal    :", journal_result["title"])
    print("Database        :", DATABASE_PATH)