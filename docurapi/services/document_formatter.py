from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from docurapi.services.file_service import clean_text

CHAPTER_PATTERN = re.compile(r"^\s*BAB\s+([IVXLCDM]+|\d+)\b", re.IGNORECASE)
SUBHEADING_PATTERN = re.compile(r"^\s*(\d+\.\d+(?:\.\d+)*)\s+(.+)$")
TABLE_CAPTION_PATTERN = re.compile(r"^\s*Tabel\s+\d+(?:\.\d+)*", re.IGNORECASE)
FIGURE_CAPTION_PATTERN = re.compile(r"^\s*(Gambar|Figure)\s+\d+(?:\.\d+)*", re.IGNORECASE)

FRONT_HEADINGS = {
    "HALAMAN JUDUL",
    "LEMBAR PERSETUJUAN",
    "LEMBAR PENGESAHAN",
    "PERNYATAAN KEASLIAN",
    "KATA PENGANTAR",
    "DAFTAR ISI",
    "DAFTAR TABEL",
    "DAFTAR GAMBAR",
    "DAFTAR LAMPIRAN",
    "ABSTRAK",
    "ABSTRACT",
    "DAFTAR PUSTAKA",
    "REFERENCES",
}

PRESETS: dict[str, dict[str, Any]] = {
    "skripsi": {
        "font": "Times New Roman",
        "font_size": 12,
        "table_font_size": 10,
        "line_spacing": 2.0,
        "first_line_indent_cm": 1.25,
        "margin_top_cm": 4,
        "margin_bottom_cm": 3,
        "margin_left_cm": 4,
        "margin_right_cm": 3,
    },
    "laporan": {
        "font": "Arial",
        "font_size": 11,
        "table_font_size": 10,
        "line_spacing": 1.5,
        "first_line_indent_cm": 1.25,
        "margin_top_cm": 3,
        "margin_bottom_cm": 3,
        "margin_left_cm": 3,
        "margin_right_cm": 3,
    },
    "jurnal": {
        "font": "Times New Roman",
        "font_size": 11,
        "table_font_size": 9,
        "line_spacing": 1.15,
        "first_line_indent_cm": 0.75,
        "margin_top_cm": 2.5,
        "margin_bottom_cm": 2.5,
        "margin_left_cm": 2.5,
        "margin_right_cm": 2.5,
    },
}


def set_run_font(run: Any, font_name: str, size_pt: float, bold: bool | None = None) -> None:
    run.font.name = font_name
    run.font.size = Pt(size_pt)

    if bold is not None:
        run.bold = bold


def apply_document_format(input_path: Path, output_path: Path, preset: str) -> dict[str, Any]:
    selected_preset = preset.lower().strip()

    if selected_preset not in PRESETS:
        selected_preset = "skripsi"

    rules = PRESETS[selected_preset]
    document = Document(input_path)

    for section in document.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(rules["margin_top_cm"])
        section.bottom_margin = Cm(rules["margin_bottom_cm"])
        section.left_margin = Cm(rules["margin_left_cm"])
        section.right_margin = Cm(rules["margin_right_cm"])

    normal_style = document.styles["Normal"]
    normal_style.font.name = rules["font"]
    normal_style.font.size = Pt(rules["font_size"])

    counters = {
        "chapters": 0,
        "subheadings": 0,
        "captions": 0,
        "body_paragraphs": 0,
        "tables": len(document.tables),
    }

    for paragraph in document.paragraphs:
        text = clean_text(paragraph.text)

        if not text:
            continue

        upper_text = text.upper()

        if CHAPTER_PATTERN.match(text) or upper_text in FRONT_HEADINGS:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(6)

            for run in paragraph.runs:
                set_run_font(run, rules["font"], 14, True)

            counters["chapters"] += 1
            continue

        if SUBHEADING_PATTERN.match(text):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(3)

            for run in paragraph.runs:
                set_run_font(run, rules["font"], rules["font_size"], True)

            counters["subheadings"] += 1
            continue

        if TABLE_CAPTION_PATTERN.match(text) or FIGURE_CAPTION_PATTERN.match(text):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)

            for run in paragraph.runs:
                set_run_font(run, rules["font"], rules["table_font_size"], True)

            counters["captions"] += 1
            continue

        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.first_line_indent = Cm(rules["first_line_indent_cm"])
        paragraph.paragraph_format.line_spacing = rules["line_spacing"]
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)

        for run in paragraph.runs:
            set_run_font(run, rules["font"], rules["font_size"])

        counters["body_paragraphs"] += 1

    for table in document.tables:
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing = 1
                    paragraph.paragraph_format.first_line_indent = Cm(0)

                    for run in paragraph.runs:
                        set_run_font(
                            run,
                            rules["font"],
                            rules["table_font_size"],
                            row_index == 0,
                        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)

    return {
        "preset": selected_preset,
        "rules": rules,
        "counters": counters,
        "output_path": str(output_path),
    }
