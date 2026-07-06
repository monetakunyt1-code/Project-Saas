from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document

from docurapi.services.file_service import clean_text


def create_journal_draft(input_path: Path, output_path: Path) -> dict[str, Any]:
    source = Document(input_path)
    target = Document()

    paragraphs = [
        clean_text(paragraph.text)
        for paragraph in source.paragraphs
        if clean_text(paragraph.text)
    ]

    full_text = "\n".join(paragraphs)

    target.add_heading("DRAFT ARTIKEL JURNAL", level=0)

    sections = {
        "Abstrak": full_text[:1200],
        "Pendahuluan": full_text[:3000],
        "Metode Penelitian": "Bagian ini perlu disesuaikan berdasarkan metode penelitian pada dokumen sumber.",
        "Hasil dan Pembahasan": "Bagian ini perlu dikembangkan berdasarkan hasil penelitian pada dokumen sumber.",
        "Kesimpulan": "Bagian ini perlu disusun ulang berdasarkan temuan utama penelitian.",
        "Daftar Pustaka": "Sesuaikan daftar pustaka dari dokumen sumber.",
    }

    for title, body in sections.items():
        target.add_heading(title, level=1)
        target.add_paragraph(body)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    target.save(output_path)

    return {
        "message": "Draft artikel jurnal awal berhasil dibuat.",
        "output_path": str(output_path),
    }
