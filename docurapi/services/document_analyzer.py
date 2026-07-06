from __future__ import annotations

import re
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document

from docurapi.services.file_service import clean_text

CHAPTER_PATTERN = re.compile(r"^\s*BAB\s+([IVXLCDM]+|\d+)\b", re.IGNORECASE)
SUBHEADING_PATTERN = re.compile(r"^\s*(\d+\.\d+(?:\.\d+)*)\s+(.+)$")
TABLE_CAPTION_PATTERN = re.compile(r"^\s*Tabel\s+\d+(?:\.\d+)*", re.IGNORECASE)
FIGURE_CAPTION_PATTERN = re.compile(r"^\s*(Gambar|Figure)\s+\d+(?:\.\d+)*", re.IGNORECASE)
APPENDIX_PATTERN = re.compile(r"^\s*Lampiran(?:\s+\d+|\s+[A-Z])?", re.IGNORECASE)


def analyze_document(file_path: Path) -> dict[str, Any]:
    document = Document(file_path)

    chapters = []
    subheadings = []
    table_captions = []
    figure_captions = []
    appendices = []
    font_names: Counter[str] = Counter()
    font_sizes: Counter[str] = Counter()
    issues = []

    non_empty_paragraphs = 0

    for index, paragraph in enumerate(document.paragraphs):
        text = clean_text(paragraph.text)

        if not text:
            continue

        non_empty_paragraphs += 1

        if CHAPTER_PATTERN.match(text):
            chapters.append({"paragraph_index": index, "text": text})

        if SUBHEADING_PATTERN.match(text):
            subheadings.append({"paragraph_index": index, "text": text})

        if TABLE_CAPTION_PATTERN.match(text):
            table_captions.append(text)

        if FIGURE_CAPTION_PATTERN.match(text):
            figure_captions.append(text)

        if APPENDIX_PATTERN.match(text):
            appendices.append(text)

        for run in paragraph.runs:
            if not clean_text(run.text):
                continue

            if run.font.name:
                font_names[run.font.name] += 1

            if run.font.size:
                font_sizes[str(round(run.font.size.pt, 2))] += 1

    if not chapters:
        issues.append(
            {
                "severity": "warning",
                "code": "chapter_not_found",
                "message": "Judul BAB belum terdeteksi.",
            }
        )

    if document.tables and not table_captions:
        issues.append(
            {
                "severity": "warning",
                "code": "table_caption_missing",
                "message": "Dokumen memiliki tabel, tetapi caption tabel belum terdeteksi.",
            }
        )

    section_data = []

    for index, section in enumerate(document.sections):
        section_data.append(
            {
                "section": index + 1,
                "top_margin_cm": round(section.top_margin.cm, 2),
                "bottom_margin_cm": round(section.bottom_margin.cm, 2),
                "left_margin_cm": round(section.left_margin.cm, 2),
                "right_margin_cm": round(section.right_margin.cm, 2),
            }
        )

    return {
        "file": {
            "name": file_path.name,
            "size_bytes": file_path.stat().st_size,
        },
        "summary": {
            "paragraphs_total": len(document.paragraphs),
            "paragraphs_non_empty": non_empty_paragraphs,
            "tables_total": len(document.tables),
            "sections_total": len(document.sections),
            "chapters_total": len(chapters),
            "subheadings_total": len(subheadings),
            "table_captions_total": len(table_captions),
            "figure_captions_total": len(figure_captions),
            "appendices_total": len(appendices),
            "issues_total": len(issues),
        },
        "structure": {
            "chapters": chapters,
            "subheadings": subheadings,
            "table_captions": table_captions,
            "figure_captions": figure_captions,
            "appendices": appendices,
        },
        "formatting": {
            "font_names": dict(font_names.most_common()),
            "font_sizes_pt": dict(font_sizes.most_common()),
            "sections": section_data,
        },
        "issues": issues,
        "analyzed_at": datetime.now(timezone.utc).isoformat(),
    }
