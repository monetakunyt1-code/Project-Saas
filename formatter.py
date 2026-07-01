import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt
from docx.text.paragraph import Paragraph

PRESETS = {
    "skripsi": dict(label="Skripsi Umum", top=4, bottom=3, left=4, right=3, font="Times New Roman", size=12, spacing=2, indent=1.27, table_size=10, page_align=WD_ALIGN_PARAGRAPH.CENTER),
    "jurnal": dict(label="Draft Jurnal Umum", top=2.54, bottom=2.54, left=2.54, right=2.54, font="Times New Roman", size=11, spacing=1, indent=.7, table_size=9, page_align=WD_ALIGN_PARAGRAPH.RIGHT),
    "laporan": dict(label="Laporan Akademik", top=3, bottom=3, left=3, right=3, font="Times New Roman", size=12, spacing=1.5, indent=1.27, table_size=10, page_align=WD_ALIGN_PARAGRAPH.CENTER),
}

CHAPTER_RE = re.compile(r"^BAB\s+(?:[IVXLCDM]+|\d+)\b", re.I)
NUMBERED_RE = re.compile(r"^(\d+(?:\.\d+){1,3})\s+\S+")
TABLE_RE = re.compile(r"^(?:TABEL|TABLE)\s+\d", re.I)
FIGURE_RE = re.compile(r"^(?:GAMBAR|FIGURE|FIG\.)\s+\d", re.I)
APPENDIX_RE = re.compile(r"^LAMPIRAN\s+(?:[A-Z]|\d+)[.:]?\s+\S+", re.I)
LIST_RE = re.compile(r"^(?:[-•●▪◦]|\(?[A-Za-z0-9]+[.)])\s+")
INDEXES = {
    "DAFTAR ISI": 'TOC \\o "1-3" \\h \\z \\u',
    "DAFTAR TABEL": "TOC \\f T \\h \\z",
    "DAFTAR GAMBAR": "TOC \\f G \\h \\z",
    "DAFTAR LAMPIRAN": "TOC \\f A \\h \\z",
}


def clean(text):
    return " ".join(text.strip().split())


def set_run(run, preset, size=None, bold=None):
    name = preset["font"]
    run.font.name = name
    run.font.size = Pt(size or preset["size"])
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for key in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{key}"), name)


def configure_styles(doc, p):
    normal = doc.styles["Normal"]
    normal.font.name = p["font"]
    normal.font.size = Pt(p["size"])
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = p["spacing"]
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    for name in ("Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        style = doc.styles[name]
        style.font.name = p["font"]
        style.font.size = Pt(p["size"])
        style.font.bold = True
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if name == "Heading 1" else WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.line_spacing = 1
        style.paragraph_format.space_before = Pt(0 if name == "Heading 1" else 6)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    caption = doc.styles["Caption"]
    caption.font.name = p["font"]
    caption.font.size = Pt(p["table_size"])
    caption.font.italic = False
    caption.paragraph_format.line_spacing = 1
    if "Judul Bagian Awal" not in [s.name for s in doc.styles]:
        front = doc.styles.add_style("Judul Bagian Awal", WD_STYLE_TYPE.PARAGRAPH)
    else:
        front = doc.styles["Judul Bagian Awal"]
    front.font.name = p["font"]
    front.font.size = Pt(p["size"])
    front.font.bold = True
    front.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    front.paragraph_format.line_spacing = 1
    front.paragraph_format.space_after = Pt(12)
    front.paragraph_format.keep_with_next = True


def page_setup(doc, p):
    for section in doc.sections:
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.top_margin, section.bottom_margin = Cm(p["top"]), Cm(p["bottom"])
        section.left_margin, section.right_margin = Cm(p["left"]), Cm(p["right"])
        section.header_distance = section.footer_distance = Cm(1.27)


def heading_level(text):
    if CHAPTER_RE.match(text):
        return 1
    match = NUMBERED_RE.match(text)
    return min(match.group(1).count(".") + 1, 4) if match else None


def add_field(paragraph, instruction, placeholder="Perbarui field di Microsoft Word."):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = f" {instruction} "
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t"); text.text = placeholder
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_tc(paragraph, text, key):
    current = " ".join(n.text or "" for n in paragraph._p.xpath(".//w:instrText"))
    if f"\\f {key}" in current:
        return
    add_field(paragraph, f'TC "{text.replace(chr(34), chr(39))}" \\f {key} \\l 1', "")


def format_paragraphs(doc, p):
    report = dict(paragraphs=0, heading_1=0, heading_2=0, heading_3=0, table_captions=0, figure_captions=0, appendix_captions=0, bibliography=0)
    in_refs = False
    chapter_count = 0
    front_titles = set(INDEXES) | {"ABSTRAK", "ABSTRACT", "DAFTAR PUSTAKA", "REFERENCES", "REFERENSI"}
    for paragraph in doc.paragraphs:
        text = clean(paragraph.text)
        if not text:
            continue
        report["paragraphs"] += 1
        upper = text.upper()
        if upper in {"DAFTAR PUSTAKA", "REFERENCES", "REFERENSI"}:
            in_refs = True
        elif upper.startswith("LAMPIRAN") or CHAPTER_RE.match(text):
            in_refs = False
        if upper in front_titles:
            paragraph.style = doc.styles["Judul Bagian Awal"]
            paragraph.paragraph_format.first_line_indent = None
            paragraph.paragraph_format.page_break_before = upper not in {"ABSTRAK", "ABSTRACT"}
            for run in paragraph.runs: set_run(run, p, bold=True)
            continue
        level = heading_level(text)
        if level:
            paragraph.style = doc.styles[f"Heading {level}"]
            paragraph.paragraph_format.first_line_indent = None
            if level == 1:
                chapter_count += 1
                paragraph.paragraph_format.page_break_before = chapter_count > 1
                report["heading_1"] += 1
            elif level == 2:
                report["heading_2"] += 1
            else:
                report["heading_3"] += 1
            for run in paragraph.runs: set_run(run, p, bold=True)
            continue
        caption_type = None
        if TABLE_RE.match(text): caption_type = ("T", "table_captions", WD_ALIGN_PARAGRAPH.LEFT)
        elif FIGURE_RE.match(text): caption_type = ("G", "figure_captions", WD_ALIGN_PARAGRAPH.CENTER)
        elif APPENDIX_RE.match(text): caption_type = ("A", "appendix_captions", WD_ALIGN_PARAGRAPH.LEFT)
        if caption_type:
            key, counter, alignment = caption_type
            paragraph.style = doc.styles["Caption"]
            paragraph.alignment = alignment
            paragraph.paragraph_format.first_line_indent = None
            add_tc(paragraph, text, key)
            report[counter] += 1
            for run in paragraph.runs: set_run(run, p, p["table_size"])
            continue
        pf = paragraph.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.space_before = pf.space_after = Pt(0)
        pf.line_spacing = p["spacing"]
        if in_refs:
            pf.left_indent, pf.first_line_indent, pf.line_spacing = Cm(1.27), Cm(-1.27), 1
            report["bibliography"] += 1
        elif LIST_RE.match(text):
            pf.first_line_indent = None
        else:
            pf.first_line_indent = Cm(p["indent"])
        for run in paragraph.runs: set_run(run, p)
    return report


def format_tables(doc, p):
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        try: table.style = "Table Grid"
        except KeyError: pass
        for row_index, row in enumerate(table.rows):
            if row_index == 0:
                trpr = row._tr.get_or_add_trPr()
                header = OxmlElement("w:tblHeader"); header.set(qn("w:val"), "true"); trpr.append(header)
            for cell in row.cells:
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                if row_index == 0:
                    tcpr = cell._tc.get_or_add_tcPr()
                    shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "D9EAF7"); tcpr.append(shade)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.first_line_indent = None
                    paragraph.paragraph_format.line_spacing = 1
                    paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
                    for run in paragraph.runs: set_run(run, p, p["table_size"], True if row_index == 0 else None)
    return len(doc.tables)


def insert_after(paragraph):
    new = OxmlElement("w:p")
    paragraph._p.addnext(new)
    return Paragraph(new, paragraph._parent)


def ensure_indexes(doc, main, enabled):
    status = {title: False for title in INDEXES}
    if not enabled:
        return status
    existing = {clean(p.text).upper(): p for p in doc.paragraphs if clean(p.text).upper() in INDEXES}
    for title, instruction in INDEXES.items():
        heading = existing.get(title)
        if heading:
            heading.paragraph_format.page_break_before = True
            next_el = heading._p.getnext()
            next_p = Paragraph(next_el, heading._parent) if next_el is not None and next_el.tag == qn("w:p") else None
            current = " ".join(n.text or "" for n in next_p._p.xpath(".//w:instrText")) if next_p else ""
            if "TOC" not in current:
                add_field(insert_after(heading), instruction)
            status[title] = True
        elif main:
            title_el = OxmlElement("w:p"); main._p.addprevious(title_el)
            title_p = Paragraph(title_el, main._parent); title_p.style = doc.styles["Judul Bagian Awal"]
            title_p.paragraph_format.page_break_before = True; title_p.add_run(title)
            field_el = OxmlElement("w:p"); main._p.addprevious(field_el)
            add_field(Paragraph(field_el, main._parent), instruction)
            status[title] = True
    return status


def has_section_break(element):
    if element is None or element.tag != qn("w:p"):
        return False
    ppr = element.find(qn("w:pPr"))
    return ppr is not None and ppr.find(qn("w:sectPr")) is not None


def insert_section_before(doc, paragraph):
    if has_section_break(paragraph._p.getprevious()):
        return
    sect = deepcopy(list(doc.sections)[-1]._sectPr)
    typ = sect.find(qn("w:type"))
    if typ is None:
        typ = OxmlElement("w:type"); sect.insert(0, typ)
    typ.set(qn("w:val"), "nextPage")
    break_p = OxmlElement("w:p"); ppr = OxmlElement("w:pPr"); ppr.append(sect); break_p.append(ppr)
    paragraph._p.addprevious(break_p)


def section_index(doc, target):
    index = 0
    for child in doc._element.body.iterchildren():
        if child is target._p: return index
        if has_section_break(child): index += 1
    return min(index, len(doc.sections) - 1)


def page_number_type(section, fmt, start=None):
    node = section._sectPr.find(qn("w:pgNumType"))
    if node is None:
        node = OxmlElement("w:pgNumType"); section._sectPr.append(node)
    node.set(qn("w:fmt"), fmt)
    if start is None: node.attrib.pop(qn("w:start"), None)
    else: node.set(qn("w:start"), str(start))


def footer_page(section, alignment, blank_first=False):
    section.footer.is_linked_to_previous = False
    paragraph = section.footer.paragraphs[0]; paragraph.clear(); paragraph.alignment = alignment
    add_field(paragraph, "PAGE", "1")
    section.different_first_page_header_footer = blank_first
    if blank_first:
        section.first_page_footer.is_linked_to_previous = False
        section.first_page_footer.paragraphs[0].clear()


def configure_pages(doc, p, main):
    if main is None:
        page_number_type(doc.sections[0], "decimal", 1)
        footer_page(doc.sections[0], p["page_align"], True)
        return {"front_sections": 0, "main_sections": len(doc.sections)}
    insert_section_before(doc, main)
    main_index = section_index(doc, main)
    sections = list(doc.sections)
    for i, section in enumerate(sections):
        if i < main_index:
            page_number_type(section, "lowerRoman", 1 if i == 0 else None)
            if i == 0: footer_page(section, p["page_align"], True)
            else: section.footer.is_linked_to_previous = True
        else:
            page_number_type(section, "decimal", 1 if i == main_index else None)
            if i == main_index: footer_page(section, p["page_align"], False)
            else: section.footer.is_linked_to_previous = True
    return {"front_sections": main_index, "main_sections": len(sections) - main_index}


def enable_update_fields(doc):
    node = doc.settings.element.find(qn("w:updateFields"))
    if node is None:
        node = OxmlElement("w:updateFields"); doc.settings.element.append(node)
    node.set(qn("w:val"), "true")


def format_document(input_path: Path, output_path: Path, preset_name: str, generate_indexes=True):
    if preset_name not in PRESETS:
        raise ValueError("Preset tidak dikenal.")
    p = PRESETS[preset_name]
    doc = Document(str(input_path))
    configure_styles(doc, p)
    page_setup(doc, p)
    main = next((x for x in doc.paragraphs if re.match(r"^BAB\s+(?:I|1)\b", clean(x.text), re.I)), None)
    r = format_paragraphs(doc, p)
    table_count = format_tables(doc, p)
    indexes = ensure_indexes(doc, main, generate_indexes)
    pages = configure_pages(doc, p, main)
    page_setup(doc, p)
    enable_update_fields(doc)
    doc.core_properties.title = f"Dokumen hasil format - {p['label']}"
    doc.core_properties.subject = "Diformat otomatis oleh DocuRapi"
    doc.save(str(output_path))
    return {
        "preset": preset_name, "preset_label": p["label"],
        "paragraphs_processed": r["paragraphs"], "heading_1_detected": r["heading_1"],
        "heading_2_detected": r["heading_2"], "heading_3_plus_detected": r["heading_3"],
        "table_captions_detected": r["table_captions"], "figure_captions_detected": r["figure_captions"],
        "appendix_captions_detected": r["appendix_captions"], "bibliography_entries_formatted": r["bibliography"],
        "tables_formatted": table_count, "indexes": indexes, "pagination": pages,
        "notes": ["Tekan Ctrl+A lalu F9 di Microsoft Word untuk memperbarui field."]
    }
